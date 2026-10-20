from fastapi import FastAPI, APIRouter, HTTPException, Request, Response, Depends
from fastapi.responses import StreamingResponse, JSONResponse
from fastapi.exceptions import RequestValidationError
from dotenv import load_dotenv
from starlette.middleware.cors import CORSMiddleware
from motor.motor_asyncio import AsyncIOMotorClient
from pymongo import ReturnDocument
from pydantic import BaseModel, Field, ConfigDict
from typing import Optional, List
import os
import io
import jwt
from bson import ObjectId
import bcrypt
import uuid
import secrets
import re
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.units import mm
from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer, PageBreak
from reportlab.graphics.shapes import Drawing, Rect, String

# Embarque les polices TTF Liberation dans tous les PDF générés pour que le rendu
# soit identique sur tous les navigateurs/téléphones (Chrome, Firefox, Opera, Mi, Samsung, iOS).
from pdf_fonts import ensure_pdf_fonts  # noqa: E402
ensure_pdf_fonts()
from datetime import datetime, timezone, timedelta
from pathlib import Path

ROOT_DIR = Path(__file__).parent
import logging  # noqa: E402
logger = logging.getLogger("golden_pirates")
load_dotenv(ROOT_DIR / '.env')

MONGO_URL = os.environ['MONGO_URL']
DB_NAME = os.environ['DB_NAME']
client = AsyncIOMotorClient(MONGO_URL)
db = client[DB_NAME]

JWT_SECRET = os.environ['JWT_SECRET']
JWT_ALGORITHM = "HS256"
ADMIN_EMAIL = os.environ.get('ADMIN_EMAIL', 'admin@goldenpirates.io').strip().lower()
ADMIN_PASSWORD = os.environ.get('ADMIN_PASSWORD', 'Capitaine1023!')
FRONTEND_URL = os.environ.get('FRONTEND_URL', 'http://localhost:3000').rstrip('/')

app = FastAPI(title="Golden Pirates — Royaume 1023")
api_router = APIRouter(prefix="/api")


def now_utc():
    return datetime.now(timezone.utc)


def safe_text(t: str) -> str:
    """Neutralise les surrogates isolés (payload JSON malformé) qui font échouer
    l'encodage UTF-8 de Mongo. Tout texte valide (chinois, russe, arabe, emoji…) passe intact."""
    return t.encode("utf-8", "ignore").decode("utf-8")


def hash_password(password: str) -> str:
    return bcrypt.hashpw(password.encode("utf-8"), bcrypt.gensalt()).decode("utf-8")


def verify_password(plain: str, hashed: str) -> bool:
    try:
        return bcrypt.checkpw(plain.encode("utf-8"), hashed.encode("utf-8"))
    except Exception:
        return False


def make_token(email: str, minutes: int) -> str:
    payload = {"sub": email, "exp": now_utc() + timedelta(minutes=minutes)}
    return jwt.encode(payload, JWT_SECRET, algorithm=JWT_ALGORITHM)


def set_auth_cookies(response: Response, email: str):
    token = make_token(email, 60 * 12)
    response.set_cookie(
        key="access_token", value=token, httponly=True, secure=True,
        samesite="none", max_age=60 * 60 * 12, path="/",
    )


def clean_doc(doc: dict) -> dict:
    doc = dict(doc)
    doc.pop("_id", None)
    return doc


async def require_admin(request: Request) -> str:
    token = request.cookies.get("access_token")
    if not token:
        auth_header = request.headers.get("Authorization", "")
        if auth_header.startswith("Bearer "):
            token = auth_header[7:]
    if not token:
        raise HTTPException(status_code=401, detail="Non authentifié")
    try:
        payload = jwt.decode(token, JWT_SECRET, algorithms=[JWT_ALGORITHM])
    except jwt.ExpiredSignatureError:
        raise HTTPException(status_code=401, detail="Session expirée")
    except jwt.InvalidTokenError:
        raise HTTPException(status_code=401, detail="Session invalide")
    email = (payload.get("sub") or "").strip().lower()
    user = await db.users.find_one({"email": email})
    if not user:
        raise HTTPException(status_code=401, detail="Compte introuvable")
    return email


# ---------------- Models ----------------

class ApplicationCreate(BaseModel):
    model_config = ConfigDict(extra="ignore")
    pseudo: str = Field(min_length=2, max_length=32)
    power: str = Field(min_length=1, max_length=16)
    troops: str = Field(min_length=2, max_length=60)
    activity: str = Field(min_length=2, max_length=60)
    experience: str = Field(default="", max_length=1200)
    discord: str = Field(default="", max_length=60)
    message: str = Field(min_length=10, max_length=2000)


class StatusUpdate(BaseModel):
    model_config = ConfigDict(extra="ignore")
    status: str
    note: Optional[str] = Field(default=None, max_length=2000)


class LoginInput(BaseModel):
    model_config = ConfigDict(extra="ignore")
    email: str
    password: str


class MessageCreate(BaseModel):
    model_config = ConfigDict(extra="ignore")
    text: str = Field(min_length=1, max_length=2000)


# ---------------- Public routes ----------------

@api_router.get("/")
async def root():
    return {"message": "Golden Pirates API — Royaume 1023", "status": "ok"}


# Cron endpoints must ack 2xx immediately; enqueue/background the actual work.
@api_router.post("/cron/keep-alive")
async def cron_keep_alive(request: Request):
    secret = os.environ.get("WEBHOOK_CRON_SECRET", "")
    auth = request.headers.get("authorization", "")
    if not secret or not secrets.compare_digest(auth, f"Bearer {secret}"):
        raise HTTPException(status_code=401, detail="Unauthorized")
    # Ce ping maintient l'app en éveil ET applique les durées de conservation RGPD (1×/heure max).
    purged = await run_retention()
    return {"ok": True, "ts": now_utc().isoformat(), "retention": purged}


# ---------------- Durées de conservation (RGPD art. 5-1-e) ----------------
# Ces durées sont celles annoncées dans les Conditions générales (/conditions) — les garder alignées.
RETENTION_IP_DAYS = 2            # IP technique anti-abus : effacée après 48 h
RETENTION_APPLICATION_DAYS = 365 # dossier + messages : 12 mois après la dernière décision
RETENTION_DONATION_DAYS = 3 * 365  # bordereaux de dons : 3 ans (preuve, remboursement, litige)
_RETENTION_STATE = {"at": None}


async def run_retention(force: bool = False):
    now = now_utc()
    last = _RETENTION_STATE["at"]
    if not force and last and (now - last) < timedelta(hours=1):
        return None
    _RETENTION_STATE["at"] = now
    res = {}
    ip_limit = (now - timedelta(days=RETENTION_IP_DAYS)).isoformat()
    r = await db.applications.update_many({"ip": {"$exists": True}, "created_at": {"$lt": ip_limit}}, {"$unset": {"ip": ""}})
    res["application_ips"] = r.modified_count
    r = await db.donators.update_many({"ip": {"$exists": True}, "created_at": {"$lt": ip_limit}}, {"$unset": {"ip": ""}})
    res["donation_ips"] = r.modified_count
    # Sanctuaire des refus : inutile au-delà du délai de 30 jours → supprimé
    rej_limit = (now - timedelta(days=REAPPLY_COOLDOWN_DAYS)).isoformat()
    r = await db.rejections.delete_many({"rejected_at": {"$lt": rej_limit}})
    res["rejections"] = r.deleted_count
    # Dossiers sans décision ni mouvement depuis 12 mois → supprimés avec leurs messages
    app_limit = (now - timedelta(days=RETENTION_APPLICATION_DAYS)).isoformat()
    old = await db.applications.find({"status_changed_at": {"$lt": app_limit}}, {"id": 1, "tracking_code": 1}).to_list(2000)
    if old:
        codes = [d["tracking_code"] for d in old if d.get("tracking_code")]
        ids = [d["id"] for d in old if d.get("id")]
        await db.messages.delete_many({"conv_code": {"$in": codes}})
        await db.peer_messages.delete_many({"$or": [{"from_id": {"$in": ids}}, {"to_id": {"$in": ids}}]})
        await db.applications.delete_many({"id": {"$in": ids}})
    res["applications"] = len(old)
    don_limit = (now - timedelta(days=RETENTION_DONATION_DAYS)).isoformat()
    r = await db.donators.delete_many({"created_at": {"$lt": don_limit}})
    res["donations"] = r.deleted_count
    # Tentatives de connexion Conseil : journal gardé 30 jours
    login_limit = (now - timedelta(days=30)).isoformat()
    r = await db.login_attempts.delete_many({"last_at": {"$lt": login_limit}})
    res["login_attempts"] = r.deleted_count
    return res


REAPPLY_COOLDOWN_DAYS = 30
IP_APPLICATIONS_PER_DAY = 6


def _pseudo_norm(pseudo: str) -> str:
    """Empreinte « large » d'un pseudo : sans accents, sans casse, lettres/chiffres seulement.
    « Jéan.Pirate », « jean_pirate », « J e a n-Pirate » → « jeanpirate » (anti-contournement du délai)."""
    import unicodedata
    txt = unicodedata.normalize("NFKD", pseudo or "")
    txt = "".join(ch for ch in txt if not unicodedata.combining(ch))
    return "".join(ch for ch in txt.lower() if ch.isalnum())


def _discord_key(discord: str) -> str:
    """Tag Discord normalisé (casse, espaces, @ initial, suffixe #0 des nouveaux pseudos)."""
    d = (discord or "").strip().lower().lstrip("@").replace(" ", "")
    if d.endswith("#0"):
        d = d[:-2]
    return d if len(d) >= 2 else ""


async def _pseudo_state(pseudo_key: str, now, pseudo_norm: str = "", discord_key: str = ""):
    """État d'un pseudo face au registre : 'free' | 'active' | 'accepted' | 'cooldown' (+ retry_at)."""
    # Garde 1 — un seul dossier actif par pseudo (pas de doublon qui noie le registre)
    if await db.applications.find_one({"pseudo_lower": pseudo_key, "status": {"$in": ["pending", "reserve"]}}, {"_id": 1}):
        return "active", None
    # Garde 2 — déjà accepté à bord : pas de nouvelle candidature
    if await db.applications.find_one({"pseudo_lower": pseudo_key, "status": "accepted"}, {"_id": 1}):
        return "accepted", None
    # Garde 3 — délai de réflexion après un refus (sanctuaire des refus, puis repli registre)
    cooldown_base = None
    # Même pseudo exact, même pseudo « déguisé » (accents/points/espaces) OU même tag Discord
    ors = [{"pseudo_lower": pseudo_key}]
    if pseudo_norm and len(pseudo_norm) >= 2:
        ors.append({"pseudo_norm": pseudo_norm})
    if discord_key:
        ors.append({"discord_key": discord_key})
    rejected_rec = await db.rejections.find_one({"$or": ors}, sort=[("rejected_at", -1)])
    if rejected_rec:
        cooldown_base = rejected_rec.get("rejected_at")
    if not cooldown_base:
        last_rejected = await db.applications.find_one(
            {"pseudo_lower": pseudo_key, "status": "rejected"},
            sort=[("created_at", -1)],
        )
        if last_rejected:
            cooldown_base = last_rejected.get("status_changed_at") or last_rejected.get("created_at")
    if cooldown_base:
        try:
            base_dt = datetime.fromisoformat(cooldown_base)
        except (TypeError, ValueError):
            base_dt = None
        if base_dt and base_dt.tzinfo is None:
            base_dt = base_dt.replace(tzinfo=timezone.utc)
        retry_at = base_dt + timedelta(days=REAPPLY_COOLDOWN_DAYS) if base_dt else None
        if retry_at and retry_at > now:
            return "cooldown", retry_at
    return "free", None


@api_router.get("/applications/check-pseudo")
async def check_pseudo(pseudo: str = ""):
    """Vérification en direct pendant la saisie (aucune donnée personnelle renvoyée)."""
    key = safe_text(pseudo or "").strip().lower()
    if len(key) < 2 or len(key) > 32:
        return {"state": "invalid"}
    state, retry_at = await _pseudo_state(key, now_utc(), _pseudo_norm(key))
    return {"state": state, "retry_at": retry_at.isoformat() if retry_at else None}


@api_router.post("/applications")
async def create_application(data: ApplicationCreate, request: Request):
    for _f in ("pseudo", "troops", "activity", "discord", "message"):
        setattr(data, _f, safe_text(getattr(data, _f) or ""))
    if data.experience:
        data.experience = safe_text(data.experience)
    pseudo_key = data.pseudo.strip().lower()
    now = now_utc()

    state, retry_at = await _pseudo_state(pseudo_key, now, _pseudo_norm(data.pseudo), _discord_key(data.discord))
    if state == "active":
        raise HTTPException(status_code=409, detail={
            "code": "duplicate_active",
            "msg": "Tu as déjà un dossier en examen ou en réserve — garde ton code de suivi et consulte l'Espace Recrue.",
        })
    if state == "accepted":
        raise HTTPException(status_code=409, detail={
            "code": "already_accepted",
            "msg": "Tu fais déjà partie de l'équipage, marin — vérifie ton dossier dans l'Espace Recrue.",
        })
    if state == "cooldown":
        raise HTTPException(status_code=429, detail={
            "code": "cooldown",
            "msg": "Dossier refusé récemment — la mer t'apprendra la patience, marin.",
            "retry_at": retry_at.isoformat(),
        })

    # Garde 4 — limite anti-flood par appareil (IP)
    ip = (request.headers.get("x-forwarded-for") or "").split(",")[0].strip()
    if not ip:
        ip = request.client.host if request.client else "unknown"
    day_ago = (now - timedelta(days=1)).isoformat()
    ip_count = await db.applications.count_documents({"ip": ip, "created_at": {"$gte": day_ago}})
    if ip_count >= IP_APPLICATIONS_PER_DAY:
        raise HTTPException(status_code=429, detail={
            "code": "ip_limit",
            "msg": "Trop de candidatures depuis cet appareil aujourd'hui — reviens demain, la mer ne bouge pas.",
        })

    doc = data.model_dump()
    doc["id"] = str(uuid.uuid4())
    doc["status"] = "pending"
    doc["note"] = ""
    doc["seen_status"] = "pending"
    doc["seen_note"] = ""
    doc["created_at"] = now.isoformat()
    doc["status_changed_at"] = now.isoformat()
    doc["pseudo_lower"] = pseudo_key
    doc["ip"] = ip
    code = None
    for _ in range(6):
        candidate = "GP-1023-" + "".join(
            secrets.choice("23456789ABCDEFGHJKMNPQRSTUVWXYZ") for _ in range(4)
        )
        if not await db.applications.find_one({"tracking_code": candidate}):
            code = candidate
            break
    if not code:
        raise HTTPException(status_code=500, detail="Génération du code impossible")
    doc["tracking_code"] = code
    await db.applications.insert_one(doc)
    return {"id": doc["id"], "tracking_code": code, "pseudo": doc["pseudo"]}


@api_router.get("/applications/track/{code}")
async def track_application(code: str):
    code = code.strip().upper()
    doc = await db.applications.find_one({"tracking_code": code})
    if not doc:
        raise HTTPException(status_code=404, detail="Aucun parchemin à ce code de suivi")
    new_status = doc["status"] != doc.get("seen_status", doc["status"])
    new_note = bool(doc.get("note")) and doc.get("note", "") != doc.get("seen_note", "")
    if new_status or new_note:
        await db.applications.update_one(
            {"tracking_code": code},
            {"$set": {"seen_status": doc["status"], "seen_note": doc.get("note", "")}},
        )
    reapply_at = None
    if doc["status"] == "rejected":
        base = doc.get("status_changed_at") or doc.get("created_at")
        try:
            base_dt = datetime.fromisoformat(base)
        except (TypeError, ValueError):
            base_dt = None
        if base_dt:
            if base_dt.tzinfo is None:
                base_dt = base_dt.replace(tzinfo=timezone.utc)
            reapply_at = (base_dt + timedelta(days=REAPPLY_COOLDOWN_DAYS)).isoformat()
    return {
        "tracking_code": doc["tracking_code"],
        "pseudo": doc["pseudo"],
        "power": doc.get("power", ""),
        "troops": doc.get("troops", ""),
        "activity": doc.get("activity", ""),
        "experience": doc.get("experience", ""),
        "discord": doc.get("discord", ""),
        "message": doc.get("message", ""),
        "peer_note": doc.get("peer_note", ""),
        "status": doc["status"],
        "note": doc.get("note", ""),
        "created_at": doc.get("created_at"),
        "reapply_at": reapply_at,
        "new_status": new_status,
        "new_note": new_note,
    }


@api_router.get("/notifications/{code}")
async def notifications(code: str):
    code = code.strip().upper()
    doc = await db.applications.find_one({"tracking_code": code}, {"status": 1, "note": 1, "seen_status": 1, "seen_note": 1})
    if not doc:
        raise HTTPException(status_code=404, detail="Aucun dossier à ce code de suivi")
    unread = await db.messages.count_documents({
        "conv_code": code, "sender": "council", "deleted": False, "read_by_recruit": False,
    })
    return {
        "unread_messages": unread,
        "new_status": doc["status"] != doc.get("seen_status", doc["status"]),
        "new_note": bool(doc.get("note")) and doc.get("note", "") != doc.get("seen_note", ""),
    }


# ---------------- Auth / Admin routes ----------------

@api_router.post("/admin/login")
async def admin_login(data: LoginInput, request: Request, response: Response):
    email = data.email.strip().lower()
    # Derrière l'ingress, request.client.host est l'IP d'un proxy interne (variable) :
    # identifier basé sur la VRAIE IP client (première entrée de X-Forwarded-For).
    ip = (request.headers.get("x-forwarded-for") or "").split(",")[0].strip()
    ident = f"{ip or (request.client.host if request.client else 'unknown')}:{email}"
    attempt = await db.login_attempts.find_one({"identifier": ident})
    if attempt and attempt.get("locked_until"):
        try:
            if datetime.fromisoformat(attempt["locked_until"]) > now_utc():
                raise HTTPException(
                    status_code=429,
                    detail="Trop de tentatives — réessayez dans 15 minutes",
                    headers={"X-Attempts-Left": "0"},
                )
            # Sceau expiré : on repart d'un compteur neuf (sinon re-verrouillage dès le 1er échec)
            await db.login_attempts.delete_one({"identifier": ident})
            attempt = None
        except ValueError:
            pass
    user = await db.users.find_one({"email": email})
    if not user or not verify_password(data.password, user.get("password_hash", "")):
        fails = (attempt.get("fails", 0) if attempt else 0) + 1
        update = {"$inc": {"fails": 1}, "$set": {"last_at": now_utc().isoformat()}}
        if fails >= 5:
            update["$set"]["locked_until"] = (now_utc() + timedelta(minutes=15)).isoformat()
        await db.login_attempts.update_one({"identifier": ident}, update, upsert=True)
        raise HTTPException(
            status_code=401,
            detail="Identifiants incorrects",
            headers={"X-Attempts-Left": str(max(0, 5 - fails))},
        )
    await db.login_attempts.delete_one({"identifier": ident})
    set_auth_cookies(response, email)
    return {"email": email, "role": user.get("role", "admin")}


@api_router.post("/auth/logout")
async def logout(response: Response):
    response.delete_cookie("access_token", path="/")
    return {"ok": True}


@api_router.get("/auth/me")
async def me(email: str = Depends(require_admin)):
    return {"email": email, "role": "admin"}


@api_router.get("/admin/applications")
async def list_applications(status: Optional[str] = None, _: str = Depends(require_admin)):
    query = {"status": status} if status else {}
    # Projection : l'IP technique anti-abus n'est jamais envoyée au navigateur du Conseil
    docs = await db.applications.find(query, {"ip": 0}).sort("created_at", -1).to_list(500)
    return [clean_doc(d) for d in docs]


@api_router.get("/admin/stats")
async def admin_stats(_: str = Depends(require_admin)):
    return {
        "total": await db.applications.count_documents({}),
        "pending": await db.applications.count_documents({"status": "pending"}),
        "accepted": await db.applications.count_documents({"status": "accepted"}),
        "reserve": await db.applications.count_documents({"status": "reserve"}),
        "rejected": await db.applications.count_documents({"status": "rejected"}),
        "donations_pending": await db.donators.count_documents({"status": "pending"}),
    }


# ---------------- Bouteille perdue : récupération du code de suivi ----------------

class RecoveryCreate(BaseModel):
    model_config = ConfigDict(extra="ignore")
    pseudo: str = Field(min_length=2, max_length=32)
    discord: Optional[str] = Field(default=None, max_length=60)
    power: Optional[str] = Field(default=None, max_length=16)


RECOVERY_PER_HOUR = 5


@api_router.post("/applications/recover")
async def recover_tracking_code(data: RecoveryCreate, request: Request):
    """La bouteille perdue : la recrue retrouve son code avec son pseudo + UNE
    clé d'identité — son Discord (semi-privé) OU sa puissance exacte (pour les
    joueurs privés de Discord par les blocages pays)."""
    now = now_utc()
    ip = (request.headers.get("x-forwarded-for") or "").split(",")[0].strip()
    if not ip:
        ip = request.client.host if request.client else "unknown"
    hour_ago = now - timedelta(hours=1)
    attempts = await db.recovery_attempts.count_documents({"ip": ip, "at": {"$gte": hour_ago}})
    if attempts >= RECOVERY_PER_HOUR:
        raise HTTPException(status_code=429, detail={
            "code": "recover_limit",
            "msg": "Trop de tentatives depuis cet appareil — réessaie dans une heure.",
        })
    await db.recovery_attempts.insert_one({"ip": ip, "at": now})

    pseudo_key = data.pseudo.strip().lower()
    docs = await db.applications.find({"pseudo_lower": pseudo_key}).sort("created_at", -1).to_list(50)

    def norm(v):
        return re.sub(r"\s+", "", (v or "")).lower()

    target_discord = norm(data.discord)
    target_power = norm(data.power)
    if not target_discord and not target_power:
        raise HTTPException(status_code=400, detail={
            "code": "recover_need_one",
            "msg": "Donne ton Discord OU ta puissance exacte pour retrouver ta bouteille.",
        })
    doc = next(
        (d for d in docs
         if (target_discord and norm(d.get("discord")) == target_discord)
         or (target_power and norm(d.get("power")) == target_power)),
        None,
    )
    if not doc:
        raise HTTPException(status_code=404, detail={
            "code": "recover_not_found",
            "msg": "Aucune bouteille au nom de ce pseudo avec ces infos — vérifie l'orthographe exacte de ta candidature.",
        })
    # Visibilité Conseil : toute récupération de code est horodatée sur le dossier.
    await db.applications.update_one(
        {"_id": doc["_id"]},
        {"$set": {"code_recovered_at": now.isoformat()}},
    )
    return {
        "tracking_code": doc["tracking_code"],
        "pseudo": doc["pseudo"],
        "status": doc["status"],
    }


# ---------------- Exports du registre (Excel / PDF) ----------------

from xml.sax.saxutils import escape

EXPORT_COLUMNS = [
    ("tracking_code", "CODE DE SUIVI"),
    ("pseudo", "PSEUDO"),
    ("power", "PUISSANCE"),
    ("troops", "TROUPES"),
    ("activity", "ACTIVITÉ"),
    ("discord", "DISCORD"),
    ("status", "STATUT"),
    ("note", "NOTE DU CONSEIL"),
    ("created_at", "DÉPOSÉE LE"),
    ("experience", "EXPÉRIENCE"),
    ("message", "MESSAGE"),
    ("peer_note", "NOTE DE MER"),
]

STATUS_I18N = {
    "fr": {"pending": "EN EXAMEN", "accepted": "ACCEPTÉ À BORD", "reserve": "EN RÉSERVE", "rejected": "REFUSÉE"},
    "en": {"pending": "UNDER REVIEW", "accepted": "ACCEPTED ABOARD", "reserve": "IN RESERVE", "rejected": "REJECTED"},
}

EXPORT_I18N = {
    "fr": {
        "syn_sheet": "Synthèse",
        "bios_sheet": "Bios des recrues",
        "reg_sheet": "Registre",
        "doc_title": "GOLDEN PIRATES — REGISTRE DES RECRUES",
        "doc_sub": "REGISTRE DES RECRUES — ROYAUME 1023 · CALL OF DRAGONS",
        "edited_on": "Édité le",
        "in_register": "candidature(s) au registre",
        "kpi": {"total": "Total", "pending": "En examen", "accepted": "Acceptées", "reserve": "En réserve", "rejected": "Refusées"},
        "read_order": "LIRE DANS L'ORDRE : « Bios des recrues » = la voix de l'équipage · « Registre » = le détail de chaque dossier, trié par décision",
        "bios_title": "BIOS DES RECRUES — LA VOIX DE L'ÉQUIPAGE",
        "recrues": "recrue(s)",
        "notes_mer": "note(s) de mer",
        "bios_headers": ["PSEUDO", "PUISSANCE", "STATUT", "NOTE DE MER", "DÉPOSÉE LE"],
        "calm": "La mer est calme — aucune candidature enregistrée pour l'instant.",
        "calm2": "Les recrues ne tarderont pas à frapper au pont.",
        "reg_banner": "REGISTRE COMPLET — ORDRE DES DÉCISIONS (EN EXAMEN → ACCEPTÉES → RÉSERVE → REFUSÉES)",
        "chart_title": "RÉPARTITION DES PUISSANCES",
        "chart_sub": "Répartition des {n} candidature(s) par tranche de puissance",
        "unreadable": "{n} puissance(s) non lisible(s) hors graphique",
        "bins": [["Moins de 6M", 0, 6], ["6M – 10M", 6, 10], ["10M – 15M", 10, 15], ["15M – 20M", 15, 20], ["20M et plus", 20, None]],
        "sections": {"pending": "EN EXAMEN — DOSSIERS À TRANCHER", "accepted": "ACCEPTÉES À BORD", "reserve": "EN RÉSERVE — LA RÉSERVE DE LA FLOTTE", "rejected": "REFUSÉES — DÉCISION RENDUE"},
        "dossiers_count": "DOSSIER(S)",
        "labels": {
            "dossier_n": "DOSSIER N° {n:02d}",
            "power": "PUISSANCE", "troops": "TROUPES", "activity": "ACTIVITÉ", "discord": "DISCORD",
            "filed": "DÉPOSÉE LE", "exp": "EXPÉRIENCE", "msg": "MESSAGE",
            "council_note": "NOTE DU CONSEIL", "sea_note": "NOTE DE MER",
        },
        "no_exp": "— aucune ligne dans le journal de bord —",
        "col": {"tracking_code": "CODE DE SUIVI", "pseudo": "PSEUDO", "power": "PUISSANCE", "troops": "TROUPES", "activity": "ACTIVITÉ", "discord": "DISCORD", "status": "STATUT", "note": "NOTE DU CONSEIL", "created_at": "DÉPOSÉE LE", "experience": "EXPÉRIENCE", "message": "MESSAGE", "peer_note": "NOTE DE MER"},
        "date_fmt": "%d/%m/%Y %H:%M",
        "file_suffix": "registre",
        "footer": "GOLDEN PIRATES — CONSEIL DES CAPITAINES · ROYAUME 1023 · CALL OF DRAGONS",
        "page_n": "Page {n}",
    },
    "en": {
        "syn_sheet": "Summary",
        "bios_sheet": "Recruits' bios",
        "reg_sheet": "Register",
        "doc_title": "GOLDEN PIRATES — RECRUITS' REGISTER",
        "doc_sub": "RECRUITS' REGISTER — KINGDOM 1023 · CALL OF DRAGONS",
        "edited_on": "Edited on",
        "in_register": "application(s) in the register",
        "kpi": {"total": "Total", "pending": "Under review", "accepted": "Accepted", "reserve": "In reserve", "rejected": "Rejected"},
        "read_order": "READ IN ORDER: “Recruits' bios” = the crew's voice · “Register” = full detail of every file, sorted by decision",
        "bios_title": "RECRUITS' BIOS — THE CREW'S VOICE",
        "recrues": "recruit(s)",
        "notes_mer": "sea note(s)",
        "bios_headers": ["PSEUDO", "POWER", "STATUS", "SEA NOTE", "SUBMITTED ON"],
        "calm": "The sea is calm — no application recorded yet.",
        "calm2": "Recruits won't be long in knocking on the deck.",
        "reg_banner": "FULL REGISTER — DECISION ORDER (UNDER REVIEW → ACCEPTED → RESERVE → REJECTED)",
        "chart_title": "POWER BREAKDOWN",
        "chart_sub": "Breakdown of {n} application(s) by power range",
        "unreadable": "{n} unreadable power value(s) kept out of the chart",
        "bins": [["Under 6M", 0, 6], ["6M – 10M", 6, 10], ["10M – 15M", 10, 15], ["15M – 20M", 15, 20], ["20M and more", 20, None]],
        "sections": {"pending": "UNDER REVIEW — FILES TO DECIDE", "accepted": "ACCEPTED ABOARD", "reserve": "IN RESERVE — THE FLEET'S BENCH", "rejected": "REJECTED — DECISION MADE"},
        "dossiers_count": "FILE(S)",
        "labels": {
            "dossier_n": "FILE N° {n:02d}",
            "power": "POWER", "troops": "TROOPS", "activity": "ACTIVITY", "discord": "DISCORD",
            "filed": "SUBMITTED ON", "exp": "EXPERIENCE", "msg": "MESSAGE",
            "council_note": "COUNCIL'S NOTE", "sea_note": "SEA NOTE",
        },
        "no_exp": "— no entry in the logbook —",
        "col": {"tracking_code": "TRACKING CODE", "pseudo": "PSEUDO", "power": "POWER", "troops": "TROOPS", "activity": "ACTIVITY", "discord": "DISCORD", "status": "STATUS", "note": "COUNCIL'S NOTE", "created_at": "SUBMITTED ON", "experience": "EXPERIENCE", "message": "MESSAGE", "peer_note": "SEA NOTE"},
        "date_fmt": "%m/%d/%Y %H:%M",
        "file_suffix": "register",
        "footer": "GOLDEN PIRATES — COUNCIL OF CAPTAINS · KINGDOM 1023 · CALL OF DRAGONS",
        "page_n": "Page {n}",
    },
}

STATUS_FR = {
    "pending": "EN EXAMEN",
    "accepted": "ACCEPTÉ À BORD",
    "reserve": "EN RÉSERVE",
    "rejected": "REFUSÉ",
}

NAVY_HEX = "#0B192C"
GOLD_HEX = "#D4AF37"
GRAY_HEX = "#5C6B7E"

STATUS_TINT = {  # fonds clairs (Excel sans #, PDF avec #)
    "pending": "FFF4CC",
    "accepted": "DCF3E8",
    "reserve": "FEF3C7",
    "rejected": "FEE2E2",
}

STATUS_HEX = {  # textes colorés
    "pending": "#B8860B",
    "accepted": "#1B7F4C",
    "reserve": "#B45309",
    "rejected": "#B91C1C",
}


def fmt_dt(v):
    if not v:
        return ""
    try:
        return datetime.fromisoformat(v).strftime("%d/%m/%Y %H:%M")
    except (TypeError, ValueError):
        return str(v)


def fmt_row(doc, lang="fr"):
    L = EXPORT_I18N[lang]
    row = []
    for key, _ in EXPORT_COLUMNS:
        v = doc.get(key, "")
        if key == "status":
            v = STATUS_I18N[lang].get(v, v)
        elif key == "created_at" and v:
            try:
                v = datetime.fromisoformat(v).strftime(L["date_fmt"])
            except (TypeError, ValueError):
                pass
        row.append("" if v is None else str(v))
    return row


@api_router.get("/admin/export/xlsx")
async def export_xlsx(lang: str = "fr", _: str = Depends(require_admin)):
    """Registre Excel — thème parchemin pirate (3 feuilles : Synthèse, Bios, Registre complet).
    Palette or/encre/parchemin cohérente avec le PDF et les factures."""
    L = EXPORT_I18N[lang if lang in ("fr", "en") else "fr"]
    docs = await db.applications.find({}).sort("created_at", -1).to_list(2000)

    # ===== Palette pirate (hex sans #) =====
    INK = "1A1612"
    GOLD = "D4AF37"
    GOLD_DARK = "B8860B"
    PARCH = "F3E6C8"
    PARCH_ALT = "EFD9A0"
    PARCH_LIGHT = "FBF4E0"
    WAX_RED = "8B1E1E"
    EMERALD = "1B7F4C"
    AMBER = "B45309"
    SEAL_TEXT = {
        "pending": AMBER,
        "accepted": EMERALD,
        "reserve": GOLD_DARK,
        "rejected": WAX_RED,
    }

    thin = Side(style="thin", color=GOLD_DARK)
    medium = Side(style="medium", color=GOLD_DARK)
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    thick_border = Border(left=medium, right=medium, top=medium, bottom=medium)

    wb = openpyxl.Workbook()

    # ============================================
    # Feuille 1 — SYNTHÈSE (parchemin pirate)
    # ============================================
    ws = wb.active
    ws.title = L["syn_sheet"]
    ws.sheet_view.showGridLines = False

    # Titre bandeau or sur fond encre
    ws.merge_cells("A1:F1")
    c = ws["A1"]
    c.value = L["doc_title"]
    c.font = Font(color=GOLD, bold=True, size=18, name="Arial")
    c.alignment = Alignment(horizontal="center", vertical="center")
    for col in "ABCDEF":
        ws[f"{col}1"].fill = PatternFill("solid", fgColor=INK)
    ws.row_dimensions[1].height = 42

    # Sous-titre or sur fond parchemin
    ws.merge_cells("A2:F2")
    c = ws["A2"]
    c.value = L["doc_sub"]
    c.font = Font(color=GOLD_DARK, bold=True, size=10, name="Arial")
    c.alignment = Alignment(horizontal="center", vertical="center")
    for col in "ABCDEF":
        ws[f"{col}2"].fill = PatternFill("solid", fgColor=PARCH_ALT)
    ws.row_dimensions[2].height = 22

    # Édité le ... italique
    ws.merge_cells("A3:F3")
    c = ws["A3"]
    c.value = f"{L['edited_on']} {fmt_dt(now_utc().isoformat())} · {len(docs)} {L['in_register']}"
    c.font = Font(color=INK, italic=True, size=9)
    c.alignment = Alignment(horizontal="center", vertical="center")
    for col in "ABCDEF":
        ws[f"{col}3"].fill = PatternFill("solid", fgColor=PARCH)
    ws.row_dimensions[3].height = 18

    # KPIs
    kpis = [
        (L["kpi"]["total"], len(docs), INK, PARCH_ALT),
        (L["kpi"]["pending"], sum(1 for d in docs if d.get("status") == "pending"), AMBER, "FDF6E3"),
        (L["kpi"]["accepted"], sum(1 for d in docs if d.get("status") == "accepted"), EMERALD, "E3F3EA"),
        (L["kpi"]["reserve"], sum(1 for d in docs if d.get("status") == "reserve"), GOLD_DARK, "F7EAC2"),
        (L["kpi"]["rejected"], sum(1 for d in docs if d.get("status") == "rejected"), WAX_RED, "F7E1E1"),
    ]
    r = 5
    for label, count, hexc, tint in kpis:
        lc = ws.cell(row=r, column=2, value=label.upper())
        lc.font = Font(bold=True, size=11, color=INK)
        lc.border = border
        lc.alignment = Alignment(horizontal="left", vertical="center", indent=1)
        lc.fill = PatternFill("solid", fgColor=PARCH_LIGHT)

        vc = ws.cell(row=r, column=3, value=count)
        vc.font = Font(bold=True, size=14, color=hexc, name="Arial")
        vc.fill = PatternFill("solid", fgColor=tint)
        vc.alignment = Alignment(horizontal="center", vertical="center")
        vc.border = border
        ws.row_dimensions[r].height = 24
        r += 1
    ws.column_dimensions["A"].width = 4
    ws.column_dimensions["B"].width = 32
    ws.column_dimensions["C"].width = 14
    for col in "DEF":
        ws.column_dimensions[col].width = 20

    # Note de lecture
    ws.merge_cells(start_row=r + 1, start_column=2, end_row=r + 1, end_column=6)
    nav = ws.cell(row=r + 1, column=2, value=L["read_order"])
    nav.font = Font(italic=True, size=9, color=INK)
    nav.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
    nav.fill = PatternFill("solid", fgColor=PARCH)
    ws.row_dimensions[r + 1].height = 32

    # ============================================
    # Feuille 2 — BIOS DES RECRUES (voix de l'équipage)
    # ============================================
    wsB = wb.create_sheet(L["bios_sheet"])
    wsB.sheet_view.showGridLines = False

    wsB.merge_cells("A1:E1")
    c = wsB["A1"]
    c.value = L["bios_title"]
    c.font = Font(color=GOLD, bold=True, size=15, name="Arial")
    c.alignment = Alignment(horizontal="center", vertical="center")
    for col in "ABCDE":
        wsB[f"{col}1"].fill = PatternFill("solid", fgColor=INK)
    wsB.row_dimensions[1].height = 36

    wsB.merge_cells("A2:E2")
    bios_count = sum(1 for d in docs if (d.get("peer_note") or "").strip())
    c = wsB["A2"]
    c.value = f"{len(docs)} {L['recrues']} · {bios_count} {L['notes_mer']} · {L['edited_on']} {fmt_dt(now_utc().isoformat())}"
    c.font = Font(color=INK, italic=True, size=9.5)
    c.alignment = Alignment(horizontal="center", vertical="center")
    for col in "ABCDE":
        wsB[f"{col}2"].fill = PatternFill("solid", fgColor=PARCH_ALT)
    wsB.row_dimensions[2].height = 22

    for i, h in enumerate(L["bios_headers"], 1):
        cell = wsB.cell(row=3, column=i, value=h)
        cell.fill = PatternFill("solid", fgColor=INK)
        cell.font = Font(color=GOLD, bold=True, size=10)
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = border
    wsB.row_dimensions[3].height = 24

    for n, d in enumerate(sorted(docs, key=lambda x: (x.get("pseudo") or "").lower()), 0):
        vals = [
            d.get("pseudo", ""),
            d.get("power", ""),
            STATUS_I18N[lang if lang in ("fr", "en") else "fr"].get(d.get("status"), d.get("status", "")),
            (d.get("peer_note") or "").strip() or "—",
            fmt_dt(d.get("created_at")) or "—",
        ]
        rn = 4 + n
        # Print-friendly : lignes blanches (pas de parchemin alterné)
        for i, v in enumerate(vals, 1):
            cell = wsB.cell(row=rn, column=i, value=v)
            cell.border = border
            cell.alignment = Alignment(vertical="top", wrap_text=(i == 4))
            if i == 1:
                cell.font = Font(bold=True, color=INK, size=10)
            else:
                cell.font = Font(color=INK, size=10)
        # Colonne statut : couleur de sceau (texte uniquement, pas de fond)
        sc = wsB.cell(row=rn, column=3)
        sc.font = Font(bold=True, color=SEAL_TEXT.get(d.get("status"), INK), size=10)
        sc.alignment = Alignment(horizontal="center", vertical="top")
    for i, w in enumerate([26, 14, 18, 60, 18], 1):
        wsB.column_dimensions[get_column_letter(i)].width = w
    wsB.freeze_panes = "A4"

    # ============================================
    # Feuille 3 — REGISTRE COMPLET
    # ============================================
    DECISION_ORDER = {"pending": 0, "accepted": 1, "reserve": 2, "rejected": 3}

    def _decision_key(d):
        ts = 0.0
        try:
            ts = datetime.fromisoformat(d.get("created_at")).timestamp() if d.get("created_at") else 0.0
        except (TypeError, ValueError):
            pass
        return (DECISION_ORDER.get(d.get("status"), 9), -ts)

    docs_ordered = sorted(docs, key=_decision_key)
    ws2 = wb.create_sheet(L["reg_sheet"])
    ws2.sheet_view.showGridLines = False

    ws2.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(EXPORT_COLUMNS))
    hc = ws2.cell(row=1, column=1, value=L["reg_banner"])
    hc.fill = PatternFill("solid", fgColor=INK)
    hc.font = Font(color=GOLD, bold=True, size=12, name="Arial")
    hc.alignment = Alignment(horizontal="center", vertical="center")
    ws2.row_dimensions[1].height = 32

    hdr_row = 2
    for i, k in enumerate([k for k, _ in EXPORT_COLUMNS], 1):
        cell = ws2.cell(row=hdr_row, column=i, value=L["col"][k])
        cell.fill = PatternFill("solid", fgColor=GOLD_DARK)
        cell.font = Font(color=INK, bold=True, size=10)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = border
    ws2.row_dimensions[hdr_row].height = 28

    status_idx = [k for k, _ in EXPORT_COLUMNS].index("status") + 1
    wrap_idx = {[k for k, _ in EXPORT_COLUMNS].index(k) + 1 for k in ("experience", "message", "note", "peer_note", "activity")}
    for n, d in enumerate(docs_ordered):
        rn = hdr_row + 1 + n
        # Print-friendly : lignes blanches (pas de parchemin alterné)
        for i, v in enumerate(fmt_row(d, lang), 1):
            cell = ws2.cell(row=rn, column=i, value=v)
            cell.border = border
            cell.alignment = Alignment(vertical="top", wrap_text=i in wrap_idx)
            cell.font = Font(color=INK, size=10)
        # Sceau de statut : texte coloré (fond blanc)
        sc = ws2.cell(row=rn, column=status_idx)
        sc.font = Font(bold=True, color=SEAL_TEXT.get(d.get("status"), INK), size=10)
        sc.alignment = Alignment(horizontal="center", vertical="top")
    if not docs:
        ws2.merge_cells(start_row=hdr_row + 1, start_column=1, end_row=hdr_row + 1, end_column=len(EXPORT_COLUMNS))
        cm = ws2.cell(row=hdr_row + 1, column=1, value=f"{L['calm']}  —  {L['calm2']}")
        cm.alignment = Alignment(horizontal="center", vertical="center")
        cm.font = Font(italic=True, color=INK, size=11)
        cm.fill = PatternFill("solid", fgColor=PARCH)
        ws2.row_dimensions[hdr_row + 1].height = 60
    for i, w in enumerate([15, 20, 12, 20, 24, 18, 15, 36, 17, 42, 52, 34], 1):
        ws2.column_dimensions[get_column_letter(i)].width = w
    ws2.freeze_panes = "A3"
    last_col = get_column_letter(len(EXPORT_COLUMNS))
    ws2.auto_filter.ref = f"A{hdr_row}:{last_col}{max(hdr_row + 1, ws2.max_row)}"

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    stamp = now_utc().strftime("%Y-%m-%d")
    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={
            "Content-Disposition": f'attachment; filename="golden-pirates-{L["file_suffix"]}-{lang}-{stamp}.xlsx"',
            "Cache-Control": "no-store",
        },
    )


def _paypal_url(amount: Optional[float] = None) -> Optional[str]:
    """Lien PayPal du coffre avec montant pré-rempli en EUR (None si non configuré)."""
    url = os.getenv("PAYPAL_DONATION_URL", "").strip()
    if not url:
        return None
    amt = None
    try:
        amt = float(amount) if amount is not None else None
    except (TypeError, ValueError):
        amt = None
    if amt is not None:
        amt = round(min(max(amt, 1.0), 10000.0), 2)
    if "paypal.me" in url:
        # 7 → "7EUR", 7.5 → "7.5EUR" (jamais "7.0") — devise forcée en euros
        amt_txt = f"{amt:.2f}".rstrip("0").rstrip(".") if amt is not None else None
        return f"{url.rstrip('/')}/{amt_txt}EUR" if amt_txt else url
    sep = "&" if "?" in url else "?"
    return f"{url}{sep}amount={amt}&currency_code=EUR" if amt else url


@api_router.get("/donations/paypal")
async def donations_paypal(amount: Optional[float] = None):
    full = _paypal_url(amount)
    if not full:
        return {"configured": False}
    return {"configured": True, "url": full}


class DonatorIn(BaseModel):
    model_config = ConfigDict(extra="ignore")
    pseudo: str = Field(min_length=2, max_length=40)
    amount: float = Field(ge=0.5, le=100000)
    email: Optional[str] = Field(default=None, max_length=120)
    currency: Optional[str] = Field(default="EUR", max_length=6)
    message: Optional[str] = Field(default=None, max_length=400)
    anonymous: bool = False
    # Date du don (ISO) — permet au Conseil de rattraper un don reçu hier/avant-hier
    # (ex. don disparu, saisie en retard). Vide = maintenant.
    created_at: Optional[str] = Field(default=None, max_length=40)


class ButinGoalIn(BaseModel):
    model_config = ConfigDict(extra="ignore")
    goal: float = Field(ge=1, le=100000)


def _generate_receipt_id() -> str:
    """Numéro de bordereau unique type GPD-1023-XXXX (4 car hex majuscules)."""
    return f"GPD-1023-{secrets.token_hex(2).upper()}"


async def _ensure_receipt_id(doc: dict) -> dict:
    """Rétro-compat : attribue un receipt_id aux dons historiques sans bordereau."""
    if doc.get("receipt_id"):
        return doc
    rid = _generate_receipt_id()
    # on retente en cas de collision (très improbable mais safe)
    for _ in range(5):
        exists = await db.donators.find_one({"receipt_id": rid}, {"_id": 1})
        if not exists:
            break
        rid = _generate_receipt_id()
    await db.donators.update_one({"_id": doc["_id"]}, {"$set": {"receipt_id": rid}})
    doc["receipt_id"] = rid
    return doc


async def get_butin_goal():
    # Objectif de butin réglé par le Conseil (réglage persistant) — 100 € par défaut.
    doc = await db.site_settings.find_one({"_id": "butin"})
    try:
        g = float(doc.get("goal", 0)) if doc else 0.0
    except (TypeError, ValueError):
        g = 0.0
    return round(g, 2) if g >= 1 else 100.0


# Un don sans champ `status` (gravé par le Conseil avant cette version) est considéré confirmé.
CONFIRMED_DON = {"status": {"$nin": ["pending", "rejected"]}}


def _don_status(d: dict) -> str:
    s = d.get("status")
    return s if s in ("pending", "rejected") else "confirmed"


@api_router.get("/donators")
async def donators_list():
    # Public — seulement les dons CONFIRMÉS par le Conseil ; jamais email/message.
    docs = await db.donators.find(CONFIRMED_DON).sort("created_at", -1).to_list(500)
    total = round(sum(float(d.get("amount", 0)) for d in docs), 2)
    return {
        "total": total,
        "goal": await get_butin_goal(),
        "count": len(docs),
        "configured": bool(os.getenv("PAYPAL_DONATION_URL", "").strip()),
        "items": [{
            "id": str(d["_id"]),
            "pseudo": "" if d.get("anonymous") else d.get("pseudo", ""),
            "anonymous": bool(d.get("anonymous")),
            "amount": d.get("amount", 0),
            "created_at": d.get("created_at"),
        } for d in docs],
    }


@api_router.get("/donators/profile/{pseudo}")
async def donator_profile(pseudo: str):
    """
    Fiche publique d'un donateur, cliquable depuis le Hall.
    - Agrège TOUS ses dons CONFIRMÉS (jamais les pending/rejected).
    - Si le pseudo correspond aussi à une candidature ACCEPTÉE (membre d'équipage),
      on expose quelques stats de jeu (NON privées) pour prouver que c'est un vrai joueur.
    - Jamais d'email, de code de suivi, de téléphone, de Discord brut, etc. (RGPD).
    """
    key = (pseudo or "").strip()
    if len(key) < 2 or len(key) > 60:
        raise HTTPException(status_code=404, detail="Pseudo introuvable.")

    # Recherche insensible à la casse mais exacte (évite les faux positifs)
    pseudo_rx = {"$regex": f"^{re.escape(key)}$", "$options": "i"}
    docs = await db.donators.find({**CONFIRMED_DON, "pseudo": pseudo_rx}).sort("created_at", -1).to_list(100)
    if not docs:
        raise HTTPException(status_code=404, detail="Aucun don confirmé pour ce pseudo.")

    # Si TOUS les dons sont anonymes → on ne révèle rien au public (respect du vœu du donateur)
    if all(d.get("anonymous") for d in docs):
        raise HTTPException(status_code=404, detail="Profil masqué par le donateur.")

    total = round(sum(float(d.get("amount", 0)) for d in docs), 2)
    count = len(docs)
    first_at = docs[-1].get("created_at")
    last_at = docs[0].get("created_at")

    gifts = []
    for d in docs:
        if d.get("anonymous"):
            continue  # on cache les dons spécifiquement marqués anonymes
        gifts.append({
            "amount": d.get("amount", 0),
            "currency": (d.get("currency") or "EUR").upper(),
            "created_at": d.get("created_at"),
            # Message volontairement limité à 160 chars côté public (le reste reste privé Conseil)
            "message": (d.get("message") or "")[:160] or None,
        })

    # Lien éventuel avec un marin accepté — on expose uniquement des données publiques.
    recruit = None
    app = await db.applications.find_one({"pseudo": pseudo_rx, "status": "accepted"})
    if app:
        recruit = {
            "status": "crew",
            "troops": app.get("troops") or None,
            "troop_level": app.get("troop_level") or None,
            "power": app.get("power") or None,  # ex: "80M" (déjà saisi en bande publique)
            "experience": app.get("experience") or None,
            "activity": app.get("activity") or None,
            "joined_at": app.get("status_changed_at") or app.get("created_at"),
        }

    canonical = next((d.get("pseudo") for d in docs if not d.get("anonymous")), key)
    return {
        "pseudo": canonical,
        "total": total,
        "count": count,
        "currency": gifts[0]["currency"] if gifts else "EUR",
        "first_at": first_at,
        "last_at": last_at,
        "gifts": gifts,
        "recruit": recruit,
    }


@api_router.get("/admin/donators")
async def donators_list_admin(_=Depends(require_admin)):
    """Pont de commandement — vue complète (email, message, devise, statut) + rétro-fit receipt_id."""
    docs = await db.donators.find({"status": {"$ne": "rejected"}}).sort("created_at", -1).to_list(500)
    items = []
    for d in docs:
        d = await _ensure_receipt_id(d)
        items.append({
            "id": str(d["_id"]),
            "pseudo": d.get("pseudo", ""),
            "amount": d.get("amount", 0),
            "currency": (d.get("currency") or "EUR").upper(),
            "email": d.get("email") or "",
            "message": d.get("message") or "",
            "provider": d.get("provider") or "paypal",
            "receipt_id": d.get("receipt_id", ""),
            "status": _don_status(d),
            "declared_by_donor": bool(d.get("declared_by_donor")),
            "anonymous": bool(d.get("anonymous")),
            "created_at": d.get("created_at"),
            "confirmed_at": d.get("confirmed_at"),
        })
    confirmed = [x for x in items if x["status"] == "confirmed"]
    total = round(sum(float(x.get("amount", 0)) for x in confirmed), 2)
    pend = [x for x in items if x["status"] == "pending"]
    pending_total = round(sum(float(x.get("amount", 0)) for x in pend), 2)
    return {"items": items, "total": total, "pending": len(pend), "pending_total": pending_total, "goal": await get_butin_goal()}


class DonationDeclareIn(BaseModel):
    model_config = ConfigDict(extra="ignore")
    pseudo: str = Field(min_length=2, max_length=40)
    amount: float = Field(ge=0.5, le=10000)
    email: Optional[str] = Field(default=None, max_length=120)
    message: Optional[str] = Field(default=None, max_length=400)
    anonymous: bool = False


@api_router.post("/donations/declare")
async def donation_declare(data: DonationDeclareIn, request: Request):
    """Le donateur grave lui-même son don après PayPal → bordereau immédiat, statut « en attente »
    jusqu'à vérification par le Conseil sur son compte PayPal."""
    ip = (request.headers.get("x-forwarded-for") or "").split(",")[0].strip()
    if not ip:
        ip = request.client.host if request.client else "unknown"
    hour_ago = (now_utc() - timedelta(hours=1)).isoformat()
    if await db.donators.count_documents({"ip": ip, "created_at": {"$gte": hour_ago}}) >= 5:
        raise HTTPException(status_code=429, detail={"code": "flood", "msg": "Trop de déclarations depuis cet appareil — patiente une heure, marin."})
    email = (data.email or "").strip().lower()[:120] or None
    if email and not re.match(r"^[^@\s]+@[^@\s]+\.[^@\s]+$", email):
        raise HTTPException(status_code=422, detail={"code": "bad_email", "msg": "Courriel invalide."})
    receipt_id = _generate_receipt_id()
    for _attempt in range(5):
        if not await db.donators.find_one({"receipt_id": receipt_id}, {"_id": 1}):
            break
        receipt_id = _generate_receipt_id()
    doc = {
        "pseudo": safe_text(data.pseudo.strip())[:40],
        "amount": round(float(data.amount), 2),
        "email": email,
        "currency": "EUR",
        "message": safe_text((data.message or "").strip())[:400] or None,
        "provider": "paypal",
        "receipt_id": receipt_id,
        "status": "pending",
        "declared_by_donor": True,
        "anonymous": bool(data.anonymous),
        "ip": ip,
        "created_at": now_utc().isoformat(),
    }
    await db.donators.insert_one(doc)
    return {
        "ok": True, "receipt_id": receipt_id, "status": "pending", "pseudo": doc["pseudo"],
        "amount": doc["amount"], "anonymous": doc["anonymous"], "created_at": doc["created_at"],
        "paypal_url": _paypal_url(doc["amount"]),
    }


@api_router.get("/donations/status/{receipt_id}")
async def donation_status(receipt_id: str):
    """Le donateur suit son don : en attente / confirmé / refusé."""
    rid = (receipt_id or "").strip().upper()[:32]
    d = await db.donators.find_one({"receipt_id": rid})
    if not d:
        raise HTTPException(status_code=404, detail={"code": "not_found", "msg": "Bordereau introuvable — vérifie le code (ex: GPD-1023-ABCD)."})
    return {
        "receipt_id": rid,
        "pseudo": d.get("pseudo", ""),
        "amount": d.get("amount", 0),
        "currency": (d.get("currency") or "EUR").upper(),
        "status": _don_status(d),
        "anonymous": bool(d.get("anonymous")),
        "created_at": d.get("created_at"),
        "confirmed_at": d.get("confirmed_at"),
        "paypal_url": _paypal_url(d.get("amount")) if _don_status(d) == "pending" else None,
    }


@api_router.post("/admin/donators/{donator_id}/confirm")
async def donator_confirm(donator_id: str, _=Depends(require_admin)):
    try:
        oid = ObjectId(donator_id)
    except Exception:
        raise HTTPException(status_code=404, detail="Bordereau introuvable.")
    r = await db.donators.update_one({"_id": oid}, {"$set": {"status": "confirmed", "confirmed_at": now_utc().isoformat()}})
    if not r.matched_count:
        raise HTTPException(status_code=404, detail="Bordereau introuvable.")
    return {"ok": True}


@api_router.post("/admin/donators/{donator_id}/reject")
async def donator_reject(donator_id: str, _=Depends(require_admin)):
    try:
        oid = ObjectId(donator_id)
    except Exception:
        raise HTTPException(status_code=404, detail="Bordereau introuvable.")
    r = await db.donators.update_one({"_id": oid}, {"$set": {"status": "rejected", "rejected_at": now_utc().isoformat()}})
    if not r.matched_count:
        raise HTTPException(status_code=404, detail="Bordereau introuvable.")
    return {"ok": True}


# ---------- Purge des données de test (fonctionne sur la base de PRODUCTION) ----------
_TEST_PSEUDO_RX = re.compile(
    r"(?i)(^(test|audit|seo[-_ ]?test|ghost|dummy|fake|debug|qa[-_]|playwright|e2e|probe|smoke|fixture|lorem|edge|msgtest|peer(alpha|beta|gamma)|vraycapitaine)|testcapitaine|pirate_?test|testpirate|test\d*$)"
)
_TEST_EMAIL_RX = re.compile(r"(?i)@(example\.(com|org|net)|test\.(com|io)|mer\.io|sea\.io|localhost|invalid)$")


async def _collect_test_data():
    apps = await db.applications.find({}, {"pseudo": 1, "discord": 1, "tracking_code": 1}).to_list(5000)
    test_apps = [a for a in apps if _TEST_PSEUDO_RX.search(str(a.get("pseudo") or "")) or _TEST_PSEUDO_RX.search(str(a.get("discord") or ""))]
    codes = [a.get("tracking_code") for a in test_apps if a.get("tracking_code")]
    dons = await db.donators.find({}, {"pseudo": 1, "email": 1, "receipt_id": 1}).to_list(5000)
    test_dons = [d for d in dons if _TEST_PSEUDO_RX.search(str(d.get("pseudo") or "")) or _TEST_EMAIL_RX.search(str(d.get("email") or ""))]
    return test_apps, codes, test_dons


@api_router.get("/admin/purge-test/preview")
async def purge_test_preview(_=Depends(require_admin)):
    test_apps, codes, test_dons = await _collect_test_data()
    msgs = await db.messages.count_documents({"conv_code": {"$in": codes}}) if codes else 0
    peers = await db.peer_messages.count_documents({"$or": [{"from": {"$in": codes}}, {"to": {"$in": codes}}, {"from": None}]})
    return {
        "applications": [{"pseudo": a.get("pseudo"), "code": a.get("tracking_code")} for a in test_apps],
        "donators": [{"pseudo": d.get("pseudo"), "receipt_id": d.get("receipt_id")} for d in test_dons],
        "messages": msgs,
        "peer_messages": peers,
    }


@api_router.post("/admin/purge-test")
async def purge_test_apply(_=Depends(require_admin)):
    test_apps, codes, test_dons = await _collect_test_data()
    res = {"applications": 0, "donators": 0, "messages": 0, "peer_messages": 0, "rejections": 0}
    if test_apps:
        res["applications"] = (await db.applications.delete_many({"_id": {"$in": [a["_id"] for a in test_apps]}})).deleted_count
        pseudos = [str(a.get("pseudo") or "").lower() for a in test_apps]
        res["rejections"] = (await db.rejections.delete_many({"pseudo_lower": {"$in": pseudos}})).deleted_count
    if codes:
        res["messages"] = (await db.messages.delete_many({"conv_code": {"$in": codes}})).deleted_count
        res["peer_messages"] = (await db.peer_messages.delete_many({"$or": [{"from": {"$in": codes}}, {"to": {"$in": codes}}]})).deleted_count
    # messages pairs orphelins (sans expéditeur valide)
    res["peer_messages"] += (await db.peer_messages.delete_many({"$or": [{"from": None}, {"from": {"$exists": False}}]})).deleted_count
    if test_dons:
        res["donators"] = (await db.donators.delete_many({"_id": {"$in": [d["_id"] for d in test_dons]}})).deleted_count
    return {"ok": True, "deleted": res}


# ---------- Vérification du carnet de la recrue (sans consommer les notifications) ----------
class RackCheckIn(BaseModel):
    model_config = ConfigDict(extra="ignore")
    codes: List[str] = Field(default_factory=list, max_length=12)


@api_router.post("/applications/rack-check")
async def rack_check(data: RackCheckIn):
    codes = [str(c).strip().upper()[:24] for c in data.codes if c][:12]
    out = {}
    if codes:
        async for doc in db.applications.find({"tracking_code": {"$in": codes}}, {"tracking_code": 1, "pseudo": 1, "status": 1, "created_at": 1}):
            out[doc["tracking_code"]] = {"exists": True, "pseudo": doc.get("pseudo", ""), "status": doc.get("status", "pending"), "created_at": doc.get("created_at")}
    for c in codes:
        out.setdefault(c, {"exists": False})
    return {"items": out}


@api_router.post("/admin/donators")
async def donators_add(data: DonatorIn, _=Depends(require_admin)):
    # Le Conseil grave un donateur après chaque don PayPal reçu, et scelle un bordereau.
    receipt_id = _generate_receipt_id()
    # unicité du bordereau (anti-collision)
    for _attempt in range(5):
        exists = await db.donators.find_one({"receipt_id": receipt_id}, {"_id": 1})
        if not exists:
            break
        receipt_id = _generate_receipt_id()
    email = (data.email or "").strip().lower()[:120] or None
    currency = (data.currency or "EUR").strip().upper()[:6]
    message = safe_text((data.message or "").strip())[:400] or None
    # Date du don : saisie du Conseil (ex. don d'hier à rattraper) ou maintenant
    chosen_at = _iso_or_none(data.created_at) if getattr(data, "created_at", None) else None
    created_iso = chosen_at or now_utc().isoformat()
    doc = {
        "pseudo": safe_text(data.pseudo.strip())[:40],
        "amount": round(float(data.amount), 2),
        "email": email,
        "currency": currency,
        "message": message,
        "provider": "paypal",
        "receipt_id": receipt_id,
        "status": "confirmed",
        "anonymous": bool(data.anonymous),
        "confirmed_at": created_iso,
        "created_at": created_iso,
    }
    await db.donators.insert_one(doc)
    return {"ok": True, "receipt_id": receipt_id}


@api_router.delete("/admin/donators/{donator_id}")
async def donators_remove(donator_id: str, _=Depends(require_admin)):
    try:
        oid = ObjectId(donator_id)
    except Exception:
        raise HTTPException(status_code=404, detail="Bordereau introuvable.")
    await db.donators.delete_one({"_id": oid})
    return {"ok": True}


@api_router.put("/admin/butin/goal")
async def butin_goal_set(data: ButinGoalIn, _=Depends(require_admin)):
    # Le Conseil règle l'objectif de la barre de progression du butin.
    goal = round(float(data.goal), 2)
    await db.site_settings.update_one({"_id": "butin"}, {"$set": {"goal": goal}}, upsert=True)
    return {"ok": True, "goal": goal}


# ---------- Calendrier des saisons (réglé par le Conseil, sans redéploiement) ----------
# s2_start      : début de la Saison 2 (fin octobre) → compte à rebours « DÉBUT DE LA SAISON 2 »
# recruit_open  : ouverture des portes de migration (fin de Saison 2) — None tant que la date
#                 n'est pas annoncée par le jeu → « date annoncée prochainement »
SEASON_DEFAULT = {"s2_start": "2026-10-31T22:59:00+00:00", "recruit_open": None}  # 31/10 23:59 heure de Paris


def _iso_or_none(v):
    if v in (None, ""):
        return None
    try:
        d = datetime.fromisoformat(str(v).replace("Z", "+00:00"))
    except ValueError:
        raise HTTPException(status_code=422, detail="Date invalide.")
    if d.tzinfo is None:
        d = d.replace(tzinfo=timezone.utc)
    return d.astimezone(timezone.utc).isoformat()


async def get_season():
    doc = await db.site_settings.find_one({"_id": "season"}) or {}
    return {
        "s2_start": doc.get("s2_start", SEASON_DEFAULT["s2_start"]) or SEASON_DEFAULT["s2_start"],
        "recruit_open": doc.get("recruit_open", SEASON_DEFAULT["recruit_open"]),
    }


class SeasonIn(BaseModel):
    s2_start: Optional[str] = None
    recruit_open: Optional[str] = None


@api_router.get("/season")
async def season_public():
    return {**(await get_season()), "now": now_utc().isoformat()}


@api_router.put("/admin/season")
async def season_set(data: SeasonIn, _=Depends(require_admin)):
    s2 = _iso_or_none(data.s2_start) or SEASON_DEFAULT["s2_start"]
    ro = _iso_or_none(data.recruit_open)
    if ro and ro <= s2:
        raise HTTPException(status_code=422, detail="L'ouverture des portes (fin de Saison 2) doit être après le début de la Saison 2.")
    await db.site_settings.update_one({"_id": "season"}, {"$set": {"s2_start": s2, "recruit_open": ro}}, upsert=True)
    return {"ok": True, "s2_start": s2, "recruit_open": ro}


# Image de partage (Discord, WhatsApp, Facebook…) qui suit le calendrier du Conseil :
# date officielle de fin de Saison 2 scellée → « RECRUTEMENT OUVERT », sinon « PRÉ-INSCRIPTIONS OUVERTES ».
_OG_DIR = Path(__file__).parent / "static" / "og"


@api_router.get("/og-image.jpg")
async def og_image():
    from fastapi.responses import FileResponse
    season = await get_season()
    name = "og-recrutement.jpg" if season.get("recruit_open") else "og-preinscription.jpg"
    return FileResponse(_OG_DIR / name, media_type="image/jpeg", headers={"Cache-Control": "public, max-age=1800"})


# ---------- Factures de don (parchemin pirate PDF / Excel, FR/EN) ----------
from donations_invoice import (
    generate_invoice_pdf,
    generate_invoice_xlsx,
    generate_donations_summary_xlsx,
    invoice_filenames,
)


def _normalize_lang(lang: str) -> str:
    return "en" if (lang or "").lower() == "en" else "fr"


async def _fetch_donator_by_id(donator_id: str) -> dict:
    try:
        oid = ObjectId(donator_id)
    except Exception:
        raise HTTPException(status_code=404, detail={"code": "not_found", "msg": "Bordereau introuvable."})
    doc = await db.donators.find_one({"_id": oid})
    if not doc:
        raise HTTPException(status_code=404, detail={"code": "not_found", "msg": "Bordereau introuvable."})
    return await _ensure_receipt_id(doc)


async def _fetch_donator_by_receipt(receipt_id: str) -> dict:
    rid = (receipt_id or "").strip().upper()
    if not rid or len(rid) > 32:
        raise HTTPException(status_code=404, detail={"code": "bad_receipt", "msg": "Bordereau introuvable."})
    doc = await db.donators.find_one({"receipt_id": rid})
    if not doc:
        raise HTTPException(status_code=404, detail={"code": "not_found", "msg": "Bordereau introuvable — vérifie le code (ex: GPD-1023-ABCD)."})
    return doc


def _invoice_response(donator: dict, fmt: str, lang: str):
    lang = _normalize_lang(lang)
    rid = str(donator.get("receipt_id") or "GPD-UNKNOWN")
    pending = _don_status(donator) == "pending"
    if fmt == "pdf":
        payload = generate_invoice_pdf(donator, lang=lang, pending=pending)
        media = "application/pdf"
    else:
        payload = generate_invoice_xlsx(donator, lang=lang, pending=pending)
        media = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    filename = invoice_filenames(rid, fmt, lang)
    headers = {
        "Content-Disposition": f'attachment; filename="{filename}"',
        "Cache-Control": "no-store",
    }
    return StreamingResponse(io.BytesIO(payload), media_type=media, headers=headers)


@api_router.get("/admin/donators/export/xlsx")
async def admin_donators_summary_xlsx(lang: str = "fr", _=Depends(require_admin)):
    """Récapitulatif Excel complet (admin)."""
    lang = _normalize_lang(lang)
    docs = await db.donators.find(CONFIRMED_DON).sort("created_at", -1).to_list(1000)
    enriched = []
    for d in docs:
        d = await _ensure_receipt_id(d)
        enriched.append({
            "receipt_id": d.get("receipt_id", ""),
            "created_at": d.get("created_at"),
            "pseudo": d.get("pseudo", ""),
            "amount": float(d.get("amount", 0)),
            "currency": (d.get("currency") or "EUR").upper(),
            "email": d.get("email") or "",
            "provider": d.get("provider") or "paypal",
            "message": d.get("message") or "",
        })
    total = round(sum(float(x.get("amount", 0)) for x in enriched), 2)
    goal = await get_butin_goal()
    payload = generate_donations_summary_xlsx(enriched, total, goal, lang=lang)
    stamp = now_utc().strftime("%Y-%m-%d")
    fname = f"GOLDEN-PIRATES-BUTIN-COMPLET-{stamp}-{lang.upper()}.xlsx"
    return StreamingResponse(
        io.BytesIO(payload),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{fname}"', "Cache-Control": "no-store"},
    )


@api_router.get("/admin/donators/{donator_id}/invoice/pdf")
async def admin_donator_invoice_pdf(donator_id: str, lang: str = "fr", _=Depends(require_admin)):
    donator = await _fetch_donator_by_id(donator_id)
    return _invoice_response(donator, "pdf", lang)


@api_router.get("/admin/donators/{donator_id}/invoice/xlsx")
async def admin_donator_invoice_xlsx(donator_id: str, lang: str = "fr", _=Depends(require_admin)):
    donator = await _fetch_donator_by_id(donator_id)
    return _invoice_response(donator, "xlsx", lang)


def _require_not_rejected(donator: dict):
    """Autorise pending + confirmed (pending aura un filigrane rouge). Bloque uniquement rejected."""
    st = _don_status(donator)
    if st == "rejected":
        raise HTTPException(status_code=409, detail={"code": "rejected", "msg": "Ce don n'a pas pu être vérifié par le Conseil. Contacte golden.pirates@outlook.fr."})


@api_router.get("/donations/invoice/pdf")
async def public_donator_invoice_pdf(receipt_id: str, lang: str = "fr"):
    """Récupération publique de la facture via bordereau receipt_id.
    - pending : facture provisoire marquée "EN ATTENTE" (filigrane rouge)
    - confirmed : facture définitive
    - rejected : bloqué (409)"""
    donator = await _fetch_donator_by_receipt(receipt_id)
    _require_not_rejected(donator)
    return _invoice_response(donator, "pdf", lang)


@api_router.get("/donations/invoice/xlsx")
async def public_donator_invoice_xlsx(receipt_id: str, lang: str = "fr"):
    donator = await _fetch_donator_by_receipt(receipt_id)
    _require_not_rejected(donator)
    return _invoice_response(donator, "xlsx", lang)


@api_router.get("/admin/export/pdf")
async def export_pdf(lang: str = "fr", _: str = Depends(require_admin)):
    """Registre PDF — parchemin pirate (100% immersif, en-tête sur chaque page, cadre or).
    Dimensions pensées pour une impression A4 propre + lecture mobile/tablette/iOS."""
    L = EXPORT_I18N[lang if lang in ("fr", "en") else "fr"]
    docs = await db.applications.find({}).sort("created_at", -1).to_list(2000)

    # ===== Couleurs parchemin (identiques aux factures) =====
    GOLD = colors.HexColor("#B8860B")
    GOLD_LIGHT = colors.HexColor("#D4AF37")
    INK = colors.HexColor("#1A1612")
    PARCH = colors.HexColor("#F3E6C8")
    PARCH_SHADE = colors.HexColor("#E6D3A1")
    WAX_RED = colors.HexColor("#8B1E1E")
    WAX_RED_LIGHT = colors.HexColor("#B91C1C")
    EMERALD = colors.HexColor("#1B7F4C")
    AMBER = colors.HexColor("#B45309")

    # Marges larges pour laisser respirer l'en-tête + le pied de page sur CHAQUE page.
    TOP_M, BOT_M, LR_M = 38 * mm, 24 * mm, 16 * mm
    buf = io.BytesIO()
    doc_pdf = SimpleDocTemplate(
        buf, pagesize=A4,
        topMargin=TOP_M, bottomMargin=BOT_M, leftMargin=LR_M, rightMargin=LR_M,
        title=L["doc_title"],
        author="Conseil des Capitaines — Golden Pirates",
    )
    W = A4[0] - 2 * LR_M

    # ===== Styles typographiques =====
    sec_style = ParagraphStyle("sec", fontName="GPSans-Bold", fontSize=12.5, leading=16, textColor=GOLD_LIGHT, alignment=1)
    kpi_num = ParagraphStyle("kn", fontName="GPSans-Bold", fontSize=18, leading=22, alignment=1)
    kpi_lab = ParagraphStyle("kl", fontName="GPSans", fontSize=7.5, leading=10, alignment=1, textColor=colors.HexColor("#6B5A3E"))
    lab = ParagraphStyle("cl", fontName="GPSans-Bold", fontSize=7, leading=9, textColor=colors.HexColor("#6B5A3E"))
    val = ParagraphStyle("cv", fontName="GPSans", fontSize=9.8, leading=13, textColor=INK)
    val_italic = ParagraphStyle("vi", fontName="GPSans-Italic", fontSize=9.5, leading=13, textColor=colors.HexColor("#3A3028"))
    pseudo_style = ParagraphStyle("ps", fontName="GPSans-Bold", fontSize=14, leading=17, textColor=INK)
    chip_pending = ParagraphStyle("cp", fontName="GPSans-Bold", fontSize=10, leading=13, alignment=1, textColor=colors.white)
    empty_style = ParagraphStyle("es", fontName="GPSans-Italic", fontSize=12, leading=17, alignment=1, textColor=colors.HexColor("#6B5A3E"))

    STATUS_COLOR = {
        "pending": AMBER,
        "accepted": EMERALD,
        "reserve": GOLD,
        "rejected": WAX_RED,
    }

    # ===== Décor permanent dessiné sur CHAQUE page (print-friendly : fond blanc + bande parchemin fine + cadre or + en-tête + filigrane + pied) =====
    def _decorate(canv, _d):
        canv.saveState()
        Wp, Hp = A4
        # 1) Fond BLANC pour impression (économie d'encre ~70%)
        canv.setFillColor(colors.white)
        canv.rect(0, 0, Wp, Hp, stroke=0, fill=1)
        # Bande parchemin fine autour du cadre (haut + bas) pour garder l'esprit pirate à l'écran
        canv.setFillColor(PARCH)
        canv.rect(0, Hp - 10 * mm, Wp, 10 * mm, stroke=0, fill=1)  # bande haute
        canv.rect(0, 0, Wp, 10 * mm, stroke=0, fill=1)              # bande basse
        # 2) Double cadre or
        canv.setStrokeColor(GOLD)
        canv.setLineWidth(2)
        canv.rect(10 * mm, 10 * mm, Wp - 20 * mm, Hp - 20 * mm, stroke=1, fill=0)
        canv.setLineWidth(0.4)
        canv.rect(13 * mm, 13 * mm, Wp - 26 * mm, Hp - 26 * mm, stroke=1, fill=0)
        # 3) Filigrane GOLDEN PIRATES diagonal (très léger pour impression)
        canv.saveState()
        canv.translate(Wp / 2, Hp / 2)
        canv.rotate(28)
        canv.setFillColor(GOLD)
        canv.setFillAlpha(0.05)
        canv.setFont("GPSans-Bold", 86)
        canv.drawCentredString(0, 0, "GOLDEN PIRATES")
        canv.setFont("GPSans-Bold", 20)
        canv.setFillAlpha(0.07)
        canv.drawCentredString(0, -48, L["doc_sub"])
        canv.restoreState()
        # 4) En-tête haut (sur CHAQUE page) — bandeau sombre or, maintenu à l'INTÉRIEUR du cadre or
        canv.setFillColor(INK)
        canv.rect(LR_M, Hp - TOP_M + 5 * mm, Wp - 2 * LR_M, 20 * mm, stroke=0, fill=1)
        canv.setStrokeColor(GOLD_LIGHT)
        canv.setLineWidth(1.2)
        canv.rect(LR_M, Hp - TOP_M + 5 * mm, Wp - 2 * LR_M, 20 * mm, stroke=1, fill=0)
        # Titre + sous-titre en blanc/or
        canv.setFillColor(GOLD_LIGHT)
        canv.setFont("GPSans-Bold", 15)
        canv.drawCentredString(Wp / 2, Hp - TOP_M + 18 * mm, "GOLDEN PIRATES — ROYAUME 1023")
        canv.setFillColor(colors.HexColor("#F8E6B8"))
        canv.setFont("GPSans", 8.3)
        canv.drawCentredString(Wp / 2, Hp - TOP_M + 10.5 * mm, L["doc_sub"])
        canv.setFont("GPSans-Italic", 7.3)
        canv.drawCentredString(
            Wp / 2, Hp - TOP_M + 6.5 * mm,
            f"{L['edited_on']} {fmt_dt(now_utc().isoformat())}  ·  {len(docs)} {L['in_register']}",
        )
        # Ancre décorative (ligne or à ras du bandeau)
        canv.setStrokeColor(GOLD)
        canv.setLineWidth(0.6)
        canv.line(LR_M + 4, Hp - TOP_M + 4 * mm, Wp - LR_M - 4, Hp - TOP_M + 4 * mm)
        # 5) Pied de page (sur CHAQUE page)
        canv.setStrokeColor(GOLD)
        canv.setLineWidth(0.6)
        canv.line(LR_M + 4, BOT_M - 8 * mm, Wp - LR_M - 4, BOT_M - 8 * mm)
        canv.setFillColor(colors.HexColor("#6B5A3E"))
        canv.setFont("GPSans-Bold", 7.5)
        canv.drawString(LR_M + 4, BOT_M - 13 * mm, L["footer"])
        canv.drawRightString(Wp - LR_M - 4, BOT_M - 13 * mm, L["page_n"].format(n=canv.getPageNumber()))
        canv.restoreState()

    # ===== CONTENU =====
    flow = []

    # --- Bloc KPI (synthèse rapide) ---
    counts = {k: sum(1 for d in docs if d.get("status") == k) for k in STATUS_I18N["fr"]}
    kpis = [
        (L["kpi"]["total"], len(docs), INK, colors.HexColor("#FBF4E0")),
        (L["kpi"]["pending"], counts["pending"], AMBER, colors.HexColor("#FEF8EA")),
        (L["kpi"]["accepted"], counts["accepted"], EMERALD, colors.HexColor("#EFF8F2")),
        (L["kpi"]["reserve"], counts["reserve"], GOLD, colors.HexColor("#FBF4D9")),
        (L["kpi"]["rejected"], counts["rejected"], WAX_RED, colors.HexColor("#FCECEC")),
    ]
    kpi_cells = []
    for i, (label, count, hexc, _tint) in enumerate(kpis):
        num_style = ParagraphStyle(f"kn{i}", parent=kpi_num, textColor=hexc)
        kpi_cells.append([Paragraph(str(count), num_style), Paragraph(label.upper(), kpi_lab)])
    kpi_table = Table([kpi_cells], colWidths=[W / 5] * 5)
    kpi_style_rules = [
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("TOPPADDING", (0, 0), (-1, -1), 10),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 10),
        ("BOX", (0, 0), (-1, -1), 0.9, GOLD),
        ("INNERGRID", (0, 0), (-1, -1), 0.4, GOLD),
    ]
    for i, (_, _, _, tint) in enumerate(kpis):
        kpi_style_rules.append(("BACKGROUND", (i, 0), (i, 0), tint))
    kpi_table.setStyle(TableStyle(kpi_style_rules))
    flow += [kpi_table, Spacer(1, 6 * mm)]

    # --- Répartition des puissances (barres or) ---
    if docs:
        bins = [(label, lo, hi) for label, lo, hi in L["bins"]]
        bin_counts = [0] * len(bins)
        unreadable = 0
        for d in docs:
            pm = parse_power_m(d.get("power", ""))
            if pm is None:
                unreadable += 1
                continue
            for i, (_, lo, hi) in enumerate(bins):
                if (hi is None and pm >= lo) or (lo <= pm < hi):
                    bin_counts[i] += 1
                    break
        chart_banner = Table([[Paragraph(L["chart_title"], sec_style)]], colWidths=[W])
        chart_banner.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, -1), INK),
            ("BOX", (0, 0), (-1, -1), 1, GOLD),
            ("TOPPADDING", (0, 0), (-1, -1), 7),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 7),
        ]))
        sub_txt = L["chart_sub"].format(n=len(docs)) + (f" · {L['unreadable'].format(n=unreadable)}" if unreadable else "")
        chart_sub = Paragraph(
            f'<para align="center"><font color="#6B5A3E" size="8"><i>{escape(sub_txt)}</i></font></para>',
            ParagraphStyle("csub", fontName="GPSans", fontSize=8),
        )
        bar_h, gap, label_w = 8.5 * mm, 5.5 * mm, 30 * mm
        chart_h = (bar_h + gap) * len(bins)
        g = Drawing(W, chart_h)
        max_c = max(bin_counts) or 1
        max_w = W - label_w - 20 * mm
        y = chart_h - bar_h
        for i, (label, _, _) in enumerate(bins):
            cnt = bin_counts[i]
            w = max(max_w * cnt / max_c, 1.5) if cnt else 1.5
            # Ombre barre
            g.add(Rect(label_w + 1, y - 1, w, bar_h, fillColor=colors.HexColor("#8B6914"), strokeColor=None))
            g.add(Rect(label_w, y, w, bar_h, fillColor=GOLD_LIGHT, strokeColor=GOLD, strokeWidth=0.4))
            g.add(String(0, y + bar_h / 2 - 3.2, label, fontName="GPSans-Bold", fontSize=8.5, fillColor=INK))
            g.add(String(label_w + w + 5, y + bar_h / 2 - 3.2, f"{cnt} · {round(cnt * 100 / len(docs))}%", fontName="GPSans-Bold", fontSize=8.5, fillColor=INK))
            y -= bar_h + gap
        flow += [chart_banner, Spacer(1, 3 * mm), chart_sub, Spacer(1, 3 * mm), g, PageBreak()]

    # --- Carte de dossier (parchemin pirate) ---
    def build_card(d, numero):
        status = d.get("status", "pending")
        wax_color = STATUS_COLOR.get(status, GOLD)
        status_label = STATUS_I18N[lang if lang in ("fr", "en") else "fr"].get(status, status)
        # En-tête carte : N° + code + chip "statut" (type sceau de cire)
        meta = Paragraph(
            f'<font color="#D4AF37" size="7.5"><b>{escape(L["labels"]["dossier_n"].format(n=numero))}</b></font>'
            f'&nbsp;&nbsp;·&nbsp;&nbsp;<font color="#E6D3A1" size="8.5">{escape(d.get("tracking_code", ""))}</font>',
            ParagraphStyle(f"hm{numero}", fontName="GPSans", leading=10),
        )
        pseudo_p = Paragraph(
            f'<font color="#F3E6C8" size="14"><b>{escape(d.get("pseudo", ""))}</b></font>',
            ParagraphStyle(f"hp{numero}", fontName="GPSans", leading=17),
        )
        chip_style = ParagraphStyle(
            f"hr{numero}", fontName="GPSans-Bold", fontSize=9.5, leading=12,
            alignment=1, textColor=colors.white,
        )
        chip = Paragraph(f'<b>{escape(status_label)}</b>', chip_style)
        rows = [
            [meta, "", "", ""],
            [pseudo_p, "", "", chip],
            [Paragraph(L["labels"]["power"], lab), Paragraph(escape(d.get("power", "")) or "—", val),
             Paragraph(L["labels"]["troops"], lab), Paragraph(escape(d.get("troops", "")) or "—", val)],
            [Paragraph(L["labels"]["activity"], lab), Paragraph(escape(d.get("activity", "")) or "—", val),
             Paragraph(L["labels"]["discord"], lab), Paragraph(escape(d.get("discord", "")) or "—", val)],
            [Paragraph(L["labels"]["filed"], lab), Paragraph(fmt_dt(d.get("created_at")) or "—", val), "", ""],
            [Paragraph(L["labels"]["exp"], lab), Paragraph(escape(d.get("experience", "")) or L["no_exp"], val_italic), "", ""],
            [Paragraph(L["labels"]["msg"], lab), Paragraph(escape(d.get("message", "")) or "—", val), "", ""],
        ]
        span_rows = [4, 5, 6]
        if d.get("note"):
            rows.append([Paragraph(L["labels"]["council_note"], lab), Paragraph(escape(d["note"]), val), "", ""])
            span_rows.append(len(rows) - 1)
        if d.get("peer_note"):
            rows.append([Paragraph(L["labels"]["sea_note"], lab), Paragraph(escape(d["peer_note"]), val_italic), "", ""])
            span_rows.append(len(rows) - 1)
        t = Table(rows, colWidths=[0.17 * W, 0.33 * W, 0.17 * W, 0.33 * W])
        style = [
            ("SPAN", (0, 0), (3, 0)),
            ("SPAN", (0, 1), (2, 1)),
            # En-tête ligne 0 et 1 : fond encre (ink)
            ("BACKGROUND", (0, 0), (3, 0), INK),
            ("BACKGROUND", (0, 1), (2, 1), INK),
            # Puce "sceau de cire" (ligne 1 col 3) : couleur selon statut
            ("BACKGROUND", (3, 1), (3, 1), wax_color),
            # Corps : fond blanc partout (print-friendly), seuls les en-têtes et le sceau sont colorés
            ("BACKGROUND", (0, 2), (0, -1), colors.HexColor("#FBF4E0")),  # labels gauche très clair
            ("BACKGROUND", (2, 2), (2, 3), colors.HexColor("#FBF4E0")),  # labels droite très clair
            ("BACKGROUND", (1, 2), (1, -1), colors.white),
            ("BACKGROUND", (3, 2), (3, 3), colors.white),
            ("TEXTCOLOR", (3, 1), (3, 1), colors.white),
            ("GRID", (0, 0), (-1, -1), 0.5, GOLD),
            ("BOX", (0, 0), (-1, -1), 1.2, GOLD),
            ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
            ("ALIGN", (3, 1), (3, 1), "CENTER"),
            ("LEFTPADDING", (0, 0), (-1, -1), 8),
            ("RIGHTPADDING", (0, 0), (-1, -1), 8),
            ("TOPPADDING", (0, 0), (-1, -1), 6),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 6),
            # Les lignes pleine largeur : fond blanc sur la cellule de valeur
            ("BACKGROUND", (1, 4), (3, 4), colors.white),
            ("BACKGROUND", (1, 5), (3, 5), colors.white),
            ("BACKGROUND", (1, 6), (3, 6), colors.white),
        ]
        for r in span_rows:
            style.append(("SPAN", (1, r), (3, r)))
            style.append(("BACKGROUND", (1, r), (3, r), colors.white))
        t.setStyle(TableStyle(style))
        return t

    if docs:
        DECISION_ORDER = {"pending": 0, "accepted": 1, "reserve": 2, "rejected": 3}

        def _decision_key(d):
            ts = 0.0
            try:
                ts = datetime.fromisoformat(d.get("created_at")).timestamp() if d.get("created_at") else 0.0
            except (TypeError, ValueError):
                pass
            return (DECISION_ORDER.get(d.get("status"), 9), -ts)

        docs_sorted = sorted(docs, key=_decision_key)
        last_status = None
        for numero, d in enumerate(docs_sorted, 1):
            st = d.get("status")
            if st != last_status:
                group_n = sum(1 for x in docs_sorted if x.get("status") == st)
                sec = Table(
                    [[Paragraph(f"{L['sections'].get(st, st)} — {group_n} {L['dossiers_count']}", sec_style)]],
                    colWidths=[W],
                )
                sec.setStyle(TableStyle([
                    ("BACKGROUND", (0, 0), (-1, -1), INK),
                    ("BOX", (0, 0), (-1, -1), 1, GOLD),
                    ("LINEBELOW", (0, 0), (-1, -1), 2, STATUS_COLOR.get(st, GOLD)),
                    ("TOPPADDING", (0, 0), (-1, -1), 9),
                    ("BOTTOMPADDING", (0, -1), (-1, -1), 9),
                ]))
                flow += [sec, Spacer(1, 4 * mm)]
                last_status = st
            flow.append(build_card(d, numero))
            flow.append(Spacer(1, 4.5 * mm))
    else:
        mer = Table([[Paragraph(f"{L['calm']}<br/><br/><i>{L['calm2']}</i>", empty_style)]], colWidths=[W])
        mer.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, -1), colors.white),
            ("BOX", (0, 0), (-1, -1), 1.5, GOLD),
            ("TOPPADDING", (0, 0), (-1, -1), 30),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 30),
        ]))
        flow.append(mer)

    doc_pdf.build(flow, onFirstPage=_decorate, onLaterPages=_decorate)
    buf.seek(0)
    stamp = now_utc().strftime("%Y-%m-%d")
    return StreamingResponse(
        buf,
        media_type="application/pdf",
        headers={
            "Content-Disposition": f'attachment; filename="golden-pirates-{L["file_suffix"]}-{lang}-{stamp}.pdf"',
            "Cache-Control": "no-store",
        },
    )


@api_router.patch("/admin/applications/{app_id}")
async def update_application(app_id: str, data: StatusUpdate, _: str = Depends(require_admin)):
    if data.status not in ("pending", "accepted", "rejected", "reserve"):
        raise HTTPException(status_code=400, detail="Statut invalide")
    current = await db.applications.find_one({"id": app_id}, {"status": 1, "pseudo_lower": 1, "pseudo": 1, "discord": 1})
    if not current:
        raise HTTPException(status_code=404, detail="Candidature introuvable")
    update = {"note": safe_text(data.note or "")}
    if current["status"] != data.status:
        update["status"] = data.status
        now_iso = now_utc().isoformat()
        update["status_changed_at"] = now_iso
        pkey = current.get("pseudo_lower") or (current.get("pseudo") or "").strip().lower()
        if pkey and data.status == "rejected":
            # Le refus est inscrit au sanctuaire : il survit au brûlage du dossier.
            await db.rejections.update_one(
                {"pseudo_lower": pkey},
                {"$set": {
                    "pseudo_lower": pkey,
                    "pseudo_norm": _pseudo_norm(current.get("pseudo") or pkey),
                    "discord_key": _discord_key(current.get("discord") or ""),
                    "rejected_at": now_iso,
                }},
                upsert=True,
            )
        elif pkey:
            # Pardon explicite : repasser le dossier en examen/accepté/réserve lève le sanctuaire.
            await db.rejections.delete_one({"pseudo_lower": pkey})
    doc = await db.applications.find_one_and_update(
        {"id": app_id},
        {"$set": update},
        return_document=ReturnDocument.AFTER,
    )
    if not doc:
        raise HTTPException(status_code=404, detail="Candidature introuvable")
    return clean_doc(doc)


@api_router.delete("/admin/applications/{app_id}")
async def delete_application(app_id: str, _: str = Depends(require_admin)):
    doc = await db.applications.find_one({"id": app_id}, {"tracking_code": 1})
    if not doc:
        raise HTTPException(status_code=404, detail="Candidature introuvable")
    await db.messages.delete_many({"conv_code": doc["tracking_code"]})
    await db.applications.delete_one({"id": app_id})
    return {"ok": True}


# ---------------- Messaging (La Bouteille à la Mer) ----------------

def clean_msg(doc: dict) -> dict:
    doc = dict(doc)
    doc.pop("_id", None)
    return doc


@api_router.get("/messages/{code}")
async def get_messages(code: str):
    code = code.strip().upper()
    if not await db.applications.find_one({"tracking_code": code}):
        raise HTTPException(status_code=404, detail="Aucun dossier à ce code de suivi")
    docs = await db.messages.find({"conv_code": code}).sort("created_at", 1).to_list(1000)
    new_ids = {d["id"] for d in docs if d["sender"] == "council" and not d.get("deleted") and not d.get("read_by_recruit")}
    if new_ids:
        await db.messages.update_many(
            {"conv_code": code, "id": {"$in": list(new_ids)}},
            {"$set": {"read_by_recruit": True}},
        )
    out = []
    for d in docs:
        if d["id"] in new_ids:
            d["was_new"] = True
        out.append(clean_msg(d))
    return out


@api_router.post("/messages/{code}")
async def send_message(code: str, data: MessageCreate):
    code = code.strip().upper()
    if not await db.applications.find_one({"tracking_code": code}):
        raise HTTPException(status_code=404, detail="Aucun dossier à ce code de suivi")
    text = safe_text(data.text).strip()
    if not text:
        raise HTTPException(status_code=400, detail="Message vide")
    minute_ago = (now_utc() - timedelta(minutes=1)).isoformat()
    recent = await db.messages.count_documents({
        "conv_code": code, "sender": "recruit", "created_at": {"$gte": minute_ago},
    })
    if recent >= 10:
        raise HTTPException(status_code=429, detail="Flood")
    doc = {
        "id": str(uuid.uuid4()),
        "conv_code": code,
        "sender": "recruit",
        "text": text[:2000],
        "created_at": now_utc().isoformat(),
        "deleted": False,
        "read_by_council": False,
    }
    await db.messages.insert_one(doc)
    return clean_msg(doc)


@api_router.delete("/messages/{code}/{message_id}")
async def delete_message(code: str, message_id: str):
    code = code.strip().upper()
    res = await db.messages.find_one_and_update(
        {"id": message_id, "conv_code": code, "sender": "recruit"},
        {"$set": {"deleted": True, "text": ""}},
        return_document=ReturnDocument.AFTER,
    )
    if not res:
        raise HTTPException(status_code=404, detail="Message introuvable")
    return clean_msg(res)


# ---------------- Tanière des recrues (messagerie recrue ↔ recrue) ----------------

class PeerMessageCreate(BaseModel):
    model_config = ConfigDict(extra="ignore")
    from_code: str = Field(min_length=4, max_length=20)
    to_app_id: str = Field(min_length=8, max_length=64)
    text: str = Field(min_length=1, max_length=1000)

def parse_power_m(text):
    """'8M', '20m+', '12,5M', '950K', '9 500 000' → millions (float) ou None si illisible."""
    t = re.sub(r"[\s\u202f']", "", (text or "")).upper().replace(",", ".")
    m = re.search(r"(\d+(?:\.\d+)?)M", t)
    if m:
        try:
            return float(m.group(1))
        except ValueError:
            return None
    m = re.search(r"(\d+(?:\.\d+)?)K", t)
    if m:
        try:
            return float(m.group(1)) / 1000
        except ValueError:
            return None
    digits = t.replace(".", "").rstrip("+")
    if re.fullmatch(r"\d{4,}", digits):
        return float(digits) / 1_000_000
    return None


class PeerNoteIn(BaseModel):
    model_config = ConfigDict(extra="ignore")
    code: str
    note: str = Field(default="", max_length=200)


@api_router.post("/peers/note")
async def peers_note(data: PeerNoteIn):
    # La recrue rédige sa note de mer (bio d'une ligne) — visible sur sa fiche de marin.
    me = (data.code or "").strip().upper()
    mine = await db.applications.find_one({"tracking_code": me})
    if not mine:
        raise HTTPException(status_code=404, detail="Aucun dossier à ce code de suivi")
    note = safe_text((data.note or "").strip())[:120]
    await db.applications.update_one(
        {"tracking_code": me},
        {"$set": {"peer_note": note, "peer_note_at": now_utc().isoformat()}},
    )
    return {"ok": True, "note": note}


@api_router.get("/peers/search")
async def peers_search(q: str, code: str, min_power: Optional[float] = None, max_power: Optional[float] = None):
    # Protégée par un code de suivi valide ; n'expose JAMAIS discord / email / code d'autrui.
    me = (code or "").strip().upper()
    mine = await db.applications.find_one({"tracking_code": me})
    if not mine:
        raise HTTPException(status_code=404, detail="Aucun dossier à ce code de suivi")
    q = (q or "").strip()
    has_range = min_power is not None or max_power is not None
    if len(q) < 2 and not has_range:
        return []
    query = {"tracking_code": {"$ne": me}}
    if len(q) >= 2:
        query["pseudo"] = {"$regex": re.escape(q), "$options": "i"}
    docs = await db.applications.find(query, collation={"locale": "en", "strength": 2}).sort("pseudo", 1).to_list(200)
    out = []
    for d in docs:
        if has_range:
            pm = parse_power_m(d.get("power", ""))
            if pm is None:
                continue
            if min_power is not None and pm < min_power:
                continue
            if max_power is not None and pm >= max_power:
                continue
        out.append({
            "app_id": d["id"],
            "pseudo": d["pseudo"],
            "power": d.get("power", ""),
            "troops": d.get("troops", ""),
            "status": d.get("status", "pending"),
            "joined": d.get("created_at"),
            "note": d.get("peer_note", ""),
        })
        if len(out) >= 12:
            break
    return out

@api_router.get("/peers/threads")
async def peers_threads(code: str):
    me = (code or "").strip().upper()
    mine = await db.applications.find_one({"tracking_code": me})
    if not mine:
        raise HTTPException(status_code=404, detail="Aucun dossier à ce code de suivi")
    my_id = mine["id"]
    docs = await db.peer_messages.find({
        "$or": [{"from_id": my_id}, {"to_id": my_id}],
        "deleted": {"$ne": True},
    }).sort("created_at", -1).to_list(1000)
    threads = {}
    for d in docs:
        from_me = d["from_id"] == my_id
        peer_id = d["to_id"] if from_me else d["from_id"]
        t = threads.setdefault(peer_id, {
            "app_id": peer_id,
            "pseudo": d["to_pseudo"] if from_me else d["from_pseudo"],
            "last": d["text"][:80],
            "last_at": d["created_at"],
            "unread": 0,
        })
        if not from_me and not d.get("read"):
            t["unread"] += 1
    return list(threads.values())

@api_router.get("/peers/thread")
async def peers_thread(code: str, peer: str):
    me = (code or "").strip().upper()
    mine = await db.applications.find_one({"tracking_code": me})
    if not mine:
        raise HTTPException(status_code=404, detail="Aucun dossier à ce code de suivi")
    other = await db.applications.find_one({"id": peer})
    if not other:
        raise HTTPException(status_code=404, detail="Recrue introuvable")
    key = "|".join(sorted([mine["id"], peer]))
    docs = await db.peer_messages.find({"thread": key, "deleted": {"$ne": True}}).sort("created_at", 1).to_list(500)
    await db.peer_messages.update_many(
        {"thread": key, "from_id": peer, "read": {"$ne": True}},
        {"$set": {"read": True}},
    )
    return {
        "peer": {"app_id": peer, "pseudo": other["pseudo"]},
        "messages": [{
            "id": d["id"],
            "from_me": d["from_id"] == mine["id"],
            "text": d["text"],
            "created_at": d["created_at"],
        } for d in docs],
    }

@api_router.post("/peers/messages")
async def peers_send(data: PeerMessageCreate):
    me = (data.from_code or "").strip().upper()
    mine = await db.applications.find_one({"tracking_code": me})
    if not mine:
        raise HTTPException(status_code=404, detail="Aucun dossier à ce code de suivi")
    other = await db.applications.find_one({"id": data.to_app_id})
    if not other:
        raise HTTPException(status_code=404, detail="Recrue introuvable")
    if other["id"] == mine["id"]:
        raise HTTPException(status_code=400, detail="Tu ne peux pas t'écrire toi-même, marin.")
    text = safe_text(data.text).strip()
    if not text:
        raise HTTPException(status_code=400, detail="Message vide")
    hour_ago = (now_utc() - timedelta(hours=1)).isoformat()
    recent = await db.peer_messages.count_documents({
        "from_id": mine["id"], "created_at": {"$gte": hour_ago},
    })
    if recent >= 20:
        raise HTTPException(status_code=429, detail="Relâche la manivelle — 20 messages par heure maximum entre recrues.")
    doc = {
        "id": str(uuid.uuid4()),
        "thread": "|".join(sorted([mine["id"], other["id"]])),
        "from_id": mine["id"],
        "to_id": other["id"],
        "from_pseudo": mine["pseudo"],
        "to_pseudo": other["pseudo"],
        "text": text[:1000],
        "created_at": now_utc().isoformat(),
        "deleted": False,
        "read": False,
    }
    await db.peer_messages.insert_one(doc)
    return {"id": doc["id"], "from_me": True, "text": doc["text"], "created_at": doc["created_at"]}


@api_router.get("/admin/conversations")
async def admin_conversations(_: str = Depends(require_admin)):
    msgs = await db.messages.find({}).sort("created_at", 1).to_list(5000)
    convs = {}
    for m in msgs:
        conv = convs.setdefault(m["conv_code"], {"last": None, "last_at": "", "unread": 0, "total": 0})
        conv["total"] += 1
        conv["last_at"] = m["created_at"]
        conv["last"] = m
        if m["sender"] == "recruit" and not m.get("deleted") and not m.get("read_by_council"):
            conv["unread"] += 1
    codes = list(convs.keys())
    pseudo_by = {}
    if codes:
        apps = await db.applications.find({"tracking_code": {"$in": codes}}).to_list(200)
        pseudo_by = {a["tracking_code"]: a["pseudo"] for a in apps}
    out = []
    for code, cv in convs.items():
        last = cv["last"] or {}
        out.append({
            "code": code,
            "pseudo": pseudo_by.get(code, code),
            "last_text": "" if last.get("deleted") else last.get("text", ""),
            "last_sender": last.get("sender", ""),
            "last_at": cv["last_at"],
            "unread": cv["unread"],
            "total": cv["total"],
        })
    out.sort(key=lambda x: x["last_at"], reverse=True)
    return out


@api_router.get("/admin/conversations/{code}")
async def admin_thread(code: str, _: str = Depends(require_admin)):
    code = code.strip().upper()
    docs = await db.messages.find({"conv_code": code}).sort("created_at", 1).to_list(1000)
    new_ids = {d["id"] for d in docs if d["sender"] == "recruit" and not d.get("deleted") and not d.get("read_by_council")}
    out = []
    for d in docs:
        if d["id"] in new_ids:
            d["was_new"] = True
        out.append(clean_msg(d))
    return out


@api_router.post("/admin/conversations/{code}")
async def admin_send_message(code: str, data: MessageCreate, _: str = Depends(require_admin)):
    code = code.strip().upper()
    if not await db.applications.find_one({"tracking_code": code}):
        raise HTTPException(status_code=404, detail="Aucun dossier à ce code de suivi")
    text = safe_text(data.text).strip()
    if not text:
        raise HTTPException(status_code=400, detail="Message vide")
    minute_ago = (now_utc() - timedelta(minutes=1)).isoformat()
    recent = await db.messages.count_documents({
        "conv_code": code, "sender": "council", "created_at": {"$gte": minute_ago},
    })
    if recent >= 10:
        raise HTTPException(status_code=429, detail="Flood")
    doc = {
        "id": str(uuid.uuid4()),
        "conv_code": code,
        "sender": "council",
        "text": text[:2000],
        "created_at": now_utc().isoformat(),
        "deleted": False,
        "read_by_council": True,
        "read_by_recruit": False,
    }
    await db.messages.insert_one(doc)
    return clean_msg(doc)


@api_router.delete("/admin/conversations/{code}/{message_id}")
async def admin_delete_message(code: str, message_id: str, _: str = Depends(require_admin)):
    code = code.strip().upper()
    res = await db.messages.find_one_and_update(
        {"id": message_id, "conv_code": code, "sender": "council"},
        {"$set": {"deleted": True, "text": ""}},
        return_document=ReturnDocument.AFTER,
    )
    if not res:
        raise HTTPException(status_code=404, detail="Message introuvable")
    return clean_msg(res)


@api_router.post("/admin/conversations/{code}/read")
async def admin_mark_read(code: str, _: str = Depends(require_admin)):
    code = code.strip().upper()
    await db.messages.update_many(
        {"conv_code": code, "sender": "recruit"},
        {"$set": {"read_by_council": True}},
    )
    return {"ok": True}


# ---------------- Startup seed ----------------

async def seed():
    await db.users.create_index("email", unique=True)
    await db.applications.create_index("tracking_code", unique=True)
    await db.applications.create_index("pseudo_lower")
    await db.applications.create_index("ip")
    await db.applications.update_many(
        {"pseudo_lower": {"$exists": False}},
        [{"$set": {
            "pseudo_lower": {"$toLower": {"$trim": {"input": "$pseudo"}}},
            "status_changed_at": "$created_at",
        }}],
    )
    # Sanctuaire des refus : le délai de re-postulation survit au brûlage.
    await db.rejections.create_index("pseudo_lower", unique=True)
    # Tentatives de récupération de code : auto-purge au bout de 2 h.
    await db.recovery_attempts.create_index("at", expireAfterSeconds=7200)
    await db.rejections.create_index("pseudo_norm")
    await db.rejections.create_index("discord_key")
    async for rej_doc in db.applications.find(
        {"status": "rejected"},
        {"pseudo": 1, "pseudo_lower": 1, "discord": 1, "status_changed_at": 1, "created_at": 1},
    ).sort("created_at", 1):
        rej_key = rej_doc.get("pseudo_lower") or (rej_doc.get("pseudo") or "").strip().lower()
        if not rej_key:
            continue
        await db.rejections.update_one(
            {"pseudo_lower": rej_key},
            {"$set": {
                "rejected_at": rej_doc.get("status_changed_at") or rej_doc.get("created_at"),
                "pseudo_norm": _pseudo_norm(rej_doc.get("pseudo") or rej_key),
                "discord_key": _discord_key(rej_doc.get("discord") or ""),
            }},
            upsert=True,
        )
    # Anciens refus (dossiers déjà brûlés) : empreinte large calculée depuis le pseudo
    async for rj in db.rejections.find({"pseudo_norm": {"$exists": False}}, {"pseudo_lower": 1}):
        await db.rejections.update_one({"_id": rj["_id"]}, {"$set": {"pseudo_norm": _pseudo_norm(rj.get("pseudo_lower") or "")}})
    await db.login_attempts.create_index("identifier")
    await db.messages.create_index([("conv_code", 1), ("created_at", 1)])
    # Rétention RGPD : JAMAIS au démarrage (aucune suppression lors d'un déploiement) —
    # uniquement via le cron planifié POST /api/cron/keep-alive (1×/heure max).
    admin = await db.users.find_one({"email": ADMIN_EMAIL})
    if admin is None:
        await db.users.insert_one({
            "email": ADMIN_EMAIL,
            "password_hash": hash_password(ADMIN_PASSWORD),
            "role": "admin",
            "created_at": now_utc().isoformat(),
        })
    elif not verify_password(ADMIN_PASSWORD, admin.get("password_hash", "")):
        await db.users.update_one(
            {"email": ADMIN_EMAIL},
            {"$set": {"password_hash": hash_password(ADMIN_PASSWORD)}},
        )


@app.on_event("startup")
async def startup():
    await seed()


app.include_router(api_router)


# Bouclier : une validation échouée (ex: caractères Unicode malformés dans le JSON)
# doit renvoyer un 422 PROPRE — jamais l'écho de l'input, qui peut planter la sérialisation.
@app.exception_handler(RequestValidationError)
async def safe_validation_handler(request: Request, exc: RequestValidationError):
    return JSONResponse(
        status_code=422,
        content={"detail": "Requête invalide — vérifie les champs envoyés."},
    )

# CORS : « * » dans CORS_ORIGINS = toutes origines (Starlette renvoie alors l'origine exacte de la
# requête quand les cookies sont utilisés, ce qui reste valide pour les navigateurs).
cors_env = os.environ.get('CORS_ORIGINS', '').strip()
if cors_env == '*':
    origins = ['*']
else:
    origins = [o.strip() for o in cors_env.split(',') if o.strip()]
    if FRONTEND_URL and FRONTEND_URL not in origins:
        origins.append(FRONTEND_URL)

app.add_middleware(
    CORSMiddleware,
    allow_origins=origins,
    allow_credentials=True,
    allow_methods=["*"],
    allow_headers=["*"],
    expose_headers=["X-Attempts-Left", "Content-Disposition"],
)


@app.on_event("shutdown")
async def shutdown_db_client():
    client.close()
